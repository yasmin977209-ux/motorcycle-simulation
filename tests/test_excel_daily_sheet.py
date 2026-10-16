"""Acceptance tests for dictionary-to-row serialization in scenario DAILY sheets."""
from __future__ import annotations

import hashlib
import json
import os
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook

from scripts.excel_builder import DATE_FMT, FINANCE_FMT, write_daily_sheet


def _normalise_source_value(header: str, value):
    if header in {"date", "Final_Close_Date"} and isinstance(value, str):
        return date.fromisoformat(value)
    return value


def test_write_daily_sheet_materializes_dictionary_values_and_preserves_none():
    payload = {
        "daily": [
            {
                "date": date(2027, 1, 1),
                "Cash": 100_000,
                "Final_Close_Date": date(2027, 1, 3),
                "Accounts_Receivable": None,
            },
            {
                "date": date(2027, 1, 2),
                "Cash": 99_000,
                "Final_Close_Date": date(2027, 1, 3),
                "Accounts_Receivable": 250,
            },
        ]
    }
    headers = list(payload["daily"][0])
    wb = Workbook()
    write_daily_sheet(wb, "FIXTURE", payload)
    ws = wb["FIXTURE_DAILY"]

    assert tuple(cell.value for cell in ws[1]) == tuple(headers)
    assert ws.max_column == len(headers)
    assert ws.max_row == len(payload["daily"]) + 1
    actual = list(ws.iter_rows(min_row=2, values_only=True))
    expected = [
        tuple(row.get(header) for header in headers)
        for row in payload["daily"]
    ]
    assert actual == expected
    assert type(actual[0][headers.index("Cash")]) is int
    assert actual[0][headers.index("Accounts_Receivable")] is None
    assert ws.cell(2, headers.index("date") + 1).is_date
    assert ws.cell(2, headers.index("date") + 1).number_format == DATE_FMT
    assert ws.cell(2, headers.index("Cash") + 1).number_format == FINANCE_FMT
    assert all(value not in headers for row in actual for value in row if isinstance(value, str))


def test_empty_daily_payload_writes_header_only():
    wb = Workbook()
    write_daily_sheet(wb, "EMPTY", {"daily": []})
    ws = wb["EMPTY_DAILY"]
    assert ws.max_column == 1
    assert ws.max_row == 1
    assert ws["A1"].value == "date"


def test_real_c030_g030_p50_daily_round_trip(tmp_path: Path):
    """Round-trip all records from the authentic, identified P50 source artifact."""
    source_path = Path(os.environ["EXCEL_DAILY_P50_JSON"])
    metadata_path = Path(os.environ["EXCEL_DAILY_SCENARIO_METADATA"])
    identity_path = Path(os.environ["EXCEL_DAILY_SOURCE_IDENTITY"])
    output_path = Path(os.environ["EXCEL_DAILY_REAL_WORKBOOK"])
    report_path = Path(os.environ["EXCEL_DAILY_ACCEPTANCE_REPORT"])

    payload = json.loads(source_path.read_text(encoding="utf-8"))
    metadata = json.loads(metadata_path.read_text(encoding="utf-8"))
    source_identity = json.loads(identity_path.read_text(encoding="utf-8"))
    rows = payload["daily"]

    assert metadata["scenario"] == "C030_G030"
    assert metadata["run_id"] == "37225915321"
    assert metadata["source_sha"] == "f3af4f06fded8b7a4e8c7730b5f2d19201a0a84b"
    assert metadata["schema_version"] == "1.0"
    assert metadata["representative_trials"]["P50"] == 156
    assert payload["summary"]["scenario_id"] == "C030_G030"
    assert payload["summary"]["trial_id"] == 156
    assert len(rows) == 1568
    headers = list(rows[0])
    assert len(headers) == 20
    assert len(set(headers)) == len(headers)
    assert all(list(row) == headers for row in rows)
    assert rows[0]["date"] == "2027-01-01"
    assert rows[-1]["date"] == "2031-04-17"
    assert all(rows[index]["date"] < rows[index + 1]["date"] for index in range(len(rows) - 1))

    wb = Workbook()
    wb.remove(wb.active)
    write_daily_sheet(wb, "C030_G030", payload)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    wb.close()

    # Reopen the physical XLSX artifact; the in-memory worksheet is not evidence.
    inspected = load_workbook(output_path, data_only=False, read_only=False)
    assert inspected.sheetnames == ["C030_G030_DAILY"]
    ws = inspected["C030_G030_DAILY"]
    assert ws.max_column == len(headers)
    assert ws.max_row == len(rows) + 1
    assert tuple(cell.value for cell in ws[1]) == tuple(headers)

    date_index = headers.index("date")
    final_close_index = headers.index("Final_Close_Date")
    all_actual_rows = list(
        ws.iter_rows(min_row=2, max_col=len(headers), values_only=True)
    )
    assert len(all_actual_rows) == len(rows)

    for row_number, (actual, source) in enumerate(zip(all_actual_rows, rows, strict=True), start=2):
        expected = tuple(_normalise_source_value(header, source.get(header)) for header in headers)
        normalised_actual = tuple(
            value.date() if isinstance(value, datetime) else value
            for value in actual
        )
        assert normalised_actual == expected, f"source/workbook mismatch at row {row_number}"
        assert list(source) == headers, f"source header/order drift at source row {row_number}"
        for col_number, (header, source_value, actual_value) in enumerate(
            zip(headers, (source.get(h) for h in headers), actual, strict=True),
            start=1,
        ):
            if header in {"date", "Final_Close_Date"}:
                assert ws.cell(row_number, col_number).is_date
                assert ws.cell(row_number, col_number).number_format == DATE_FMT
            elif type(source_value) is int:
                assert type(actual_value) is int, (
                    f"integer type changed for {header} at row {row_number}"
                )
                assert ws.cell(row_number, col_number).number_format == FINANCE_FMT

    first_values = tuple(
        value.date() if isinstance(value, datetime) else value
        for value in all_actual_rows[0]
    )
    last_values = tuple(
        value.date() if isinstance(value, datetime) else value
        for value in all_actual_rows[-1]
    )
    assert first_values == tuple(_normalise_source_value(h, rows[0].get(h)) for h in headers)
    assert last_values == tuple(_normalise_source_value(h, rows[-1].get(h)) for h in headers)
    assert first_values[headers.index("Accounts_Receivable")] == 0
    assert first_values[headers.index("Cash")] == 0
    assert first_values[headers.index("Cumulative_Project_Profit")] == -100000
    assert all(
        ws.cell(row_number, date_index + 1).is_date
        and ws.cell(row_number, final_close_index + 1).is_date
        for row_number in range(2, ws.max_row + 1)
    )

    formula_cells = []
    error_cells = []
    for cell_row in ws.iter_rows():
        for cell in cell_row:
            if cell.data_type == "f":
                formula_cells.append(cell.coordinate)
            if cell.data_type == "e":
                error_cells.append(cell.coordinate)
    assert formula_cells == []
    assert error_cells == []

    output_digest = hashlib.sha256(output_path.read_bytes()).hexdigest()
    output_bytes = output_path.stat().st_size
    report = {
        "status": "PASS_REAL_P50_DAILY_ROUND_TRIP",
        "acceptance_scope": "C030_G030 P50 DAILY sheet serializer round-trip only; not full 185-sheet workbook or final delivery acceptance",
        "scenario": metadata["scenario"],
        "p50_trial_id": payload["summary"]["trial_id"],
        "source_run_id": metadata["run_id"],
        "source_sha": metadata["source_sha"],
        "schema_version": metadata["schema_version"],
        "source_artifacts": source_identity,
        "daily_row_count": len(rows),
        "header_count": len(headers),
        "header_order": headers,
        "worksheet": ws.title,
        "worksheet_dimensions": ws.calculate_dimension(),
        "first_date": rows[0]["date"],
        "last_date": rows[-1]["date"],
        "first_excel_data_row": list(first_values),
        "last_excel_data_row": list(last_values),
        "formula_cell_count": len(formula_cells),
        "excel_error_cell_count": len(error_cells),
        "workbook_path": output_path.name,
        "workbook_bytes": output_bytes,
        "workbook_sha256": output_digest,
        "tested_row_comparisons": len(rows),
        "tested_cell_comparisons": len(rows) * len(headers),
    }
    report_path.parent.mkdir(parents=True, exist_ok=True)
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    inspected.close()
    assert output_bytes > 0
    assert len(output_digest) == 64
