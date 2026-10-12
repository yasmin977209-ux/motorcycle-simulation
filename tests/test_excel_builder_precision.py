from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook

from scripts.excel_builder import save_workbook_preserving_stats_precision


def test_stats_precision_round_trips_exactly_for_gate_cell(tmp_path: Path) -> None:
    expected = float("13185720.000000002")
    assert float("%.16g" % expected) != expected

    workbook = Workbook()
    ws = workbook.active
    ws.title = "C030_G030_STATS"
    ws.append(["metric", "Mean", "Median", "Std", "Min", "Max", "P5", "P10", "P25", "P50", "P75", "P90", "P95"])
    ws.append(["Final_Cash", 1, 1, 0, 0, 20_000_000, 1, 1, 1, 1, 1, expected, 1])
    ws["L2"].number_format = "#,##0"

    other = workbook.create_sheet("CONSTANTS")
    other["A1"] = "unrelated"
    other["B1"] = expected

    output = tmp_path / "precision.xlsx"
    result = save_workbook_preserving_stats_precision(workbook, output)

    assert result["target_cells_verified"] == 1
    assert result["cells_patched"] == 1
    assert result["patched_cells"]
    assert result["numeric_cells_only"] is True
    assert result["uses_tolerance_or_rounding"] is False

    saved = load_workbook(output, data_only=True, read_only=False)
    gate_cell = saved["C030_G030_STATS"]["L2"]
    assert gate_cell.data_type == "n"
    assert type(gate_cell.value) is float
    assert gate_cell.value == expected
    assert gate_cell.number_format == "#,##0"
    assert saved["CONSTANTS"]["B1"].value == float("%.16g" % expected)
    saved.close()
