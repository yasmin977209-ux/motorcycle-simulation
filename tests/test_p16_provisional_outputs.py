from __future__ import annotations

from openpyxl import Workbook, load_workbook

from scripts import excel_builder


def _assignments():
    result = {}
    for scenario in excel_builder.EXPECTED_SCENARIOS:
        result[scenario] = {
            "P10": 1,
            "P50": 1,
            "P90": 1,
            "LossCase": 1,
            "MaxCase": 60 if scenario == "C050_G000" else 1,
        }
    return result


def _representative_payload(scenario: str, trial_id: int):
    return {
        "summary": {
            "scenario_id": scenario,
            "trial_id": trial_id,
            "final_close_date": "2033-02-12",
            "final_net_project_equity": 100,
            "final_cash": 90,
            "cumulative_project_profit": 80,
            "termination_count": 1,
            "secondary_cycle_count": 0,
            "owned_bikes": 2,
            "held_assets": 1,
            "total_operating_revenue": 200,
            "bad_debt": 10,
            "guarantee_recovered": 0,
        },
        "daily": [{"trial_id": trial_id, "date": "2033-02-12"}],
    }


def test_select_maxcase_uses_maximum_positive_integer_and_smallest_id_tie():
    from scripts.p16_produce_deliverables_v4 import select_maxcase_row

    rows = [
        {"trial_id": 17, "final_net_project_equity": 900},
        {"trial_id": 4, "final_net_project_equity": 900},
        {"trial_id": 2, "final_net_project_equity": 899},
    ]
    assert select_maxcase_row(rows)["trial_id"] == 4
    assert select_maxcase_row(rows)["final_net_project_equity"] == 900


def test_select_maxcase_fails_closed_without_positive_equity():
    from scripts.p16_produce_deliverables_v4 import select_maxcase_row

    rows = [
        {"trial_id": 1, "final_net_project_equity": 0},
        {"trial_id": 2, "final_net_project_equity": -1},
    ]
    try:
        select_maxcase_row(rows)
    except RuntimeError as exc:
        assert "UNDEFINED_BUSINESS_RULE" in str(exc)
    else:
        raise AssertionError("expected fail-closed result when no positive MaxCase exists")


def test_representative_cases_mark_only_missing_maxcase_payload_and_bind_loss_json():
    wb = Workbook()
    wb.active.title = "README"
    assignments = _assignments()
    payloads = {
        scenario: {
            role: _representative_payload(scenario, assignments[scenario][role])
            for role in ("P10", "P50", "P90", "LossCase")
        }
        for scenario in excel_builder.EXPECTED_SCENARIOS
    }
    labels = {"P10": "P10", "P50": "P50", "P90": "P90", "LossCase": "Loss", "MaxCase": "MaxCase"}
    excel_builder.write_representatives(wb, {}, payloads, assignments, labels)
    ws = wb["REPRESENTATIVE_CASES"]
    headers = [c.value for c in ws[1]]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    maxcase_rows = [row for row in rows if row[1] == "MaxCase" and row[2] == "PENDING_DATA_SOURCE"]
    loss_rows = [row for row in rows if row[1] == "LossCase" and row[2] == "SUMMARY"]
    assert len(maxcase_rows) == 25
    assert sum(row[3] == 60 for row in maxcase_rows if row[0] == "C050_G000") == 1
    source_index = headers.index("payload_source")
    assert all(row[source_index] == "representatives/Loss.json" for row in loss_rows)
    maxcase_row_num = next(i for i in range(2, ws.max_row + 1) if ws.cell(i, 1).value == "C050_G000" and ws.cell(i, 2).value == "MaxCase")
    assert ws.cell(maxcase_row_num, 1).comment.text == (
        "MaxCase trial_id=60; daily payload missing from V3 — see R4-V2.1 §5.4"
    )
    assert ws.cell(maxcase_row_num, headers.index("Partner1_Reinvestment_Balance") + 1).value == "PENDING_DATA_SOURCE"
    assert ws.cell(maxcase_row_num, headers.index("Partner2_Reinvestment_Balance") + 1).value == "PENDING_DATA_SOURCE"
    wb.close()


def test_pending_maxcase_warning_is_not_a_workbook_error():
    from scripts.stage13_assemble import assemble_workbook, expected_sheet_names

    scenarios = list(excel_builder.EXPECTED_SCENARIOS)
    roles = ("P10", "P50", "P90", "LossCase", "MaxCase")
    assignments = {s: {r: 1 for r in roles} for s in scenarios}
    trial_rows = {s: [{"trial_id": 1}] for s in scenarios}
    payloads = {s: {r: _representative_payload(s, 1) for r in roles if r != "MaxCase"} for s in scenarios}

    def global_writer(name):
        def write(wb, *args, **kwargs):
            if name == "README":
                ws = wb.active
                ws.title = "README"
                ws.append(["Excel sheets", 1])
                ws.append(["Excel error count", 0])
            else:
                wb.create_sheet(name)
        return write

    from types import SimpleNamespace

    def write_stats(wb, scenario, *args, **kwargs):
        wb.create_sheet(f"{scenario}_STATS")

    def write_details(wb, scenario, *args, **kwargs):
        for suffix in excel_builder.SUFFIXES[1:]:
            wb.create_sheet(f"{scenario}{suffix}")

    fake_builder = SimpleNamespace(
        EXPECTED_SCENARIOS=scenarios,
        GLOBAL_SHEETS=excel_builder.GLOBAL_SHEETS,
        SUFFIXES=excel_builder.SUFFIXES,
        write_readme=global_writer("README"),
        write_constants=global_writer("CONSTANTS"),
        build_scenario_summary=lambda _rows: {},
        write_scenario_matrix=global_writer("SCENARIO_MATRIX"),
        write_heatmap=global_writer("HEATMAP"),
        write_sensitivity=global_writer("SENSITIVITY"),
        write_validation=global_writer("VALIDATION_SUMMARY"),
        write_errors=global_writer("ERROR_LOG"),
        write_final_closure=global_writer("FINAL_CLOSURE_SUMMARY"),
        write_daily_distribution=global_writer("DAILY_DISTRIBUTION"),
        write_representatives=global_writer("REPRESENTATIVE_CASES"),
        write_stats_sheet=write_stats,
        write_scenario_detail=write_details,
        save_workbook_preserving_stats_precision=lambda wb, path: (wb.save(path) or {"synthetic": True}),
    )
    data = {
        "manifest": {"total_capital": 1},
        "trials_by_scenario": trial_rows,
        "representatives_by_scenario": payloads,
        "representative_trial_ids": assignments,
        "validation": {"status": "synthetic", "tests": []},
        "errors": [],
        "dataset_root": ".",
        "pending_roles": ["MaxCase"],
    }
    mapping = {"P10": "P10", "P50": "P50", "P90": "P90", "LossCase": "Loss", "MaxCase": "MaxCase"}
    result = assemble_workbook(
        builder=fake_builder,
        data=data,
        output_path="/tmp/p16-warning-contract-test.xlsx",
        role_to_payload_labels=mapping,
        mode="synthetic-test-only",
    )
    assert result["success"] is True
    assert result["excel_errors"] == 0
    assert any("MAXCASE_PAYLOAD_PENDING_DATA_SOURCE" in item for item in result["warnings"])
