from __future__ import annotations

from datetime import date
from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import Workbook, load_workbook

from scripts import excel_builder

from scripts.stage13_assemble import (
    REQUIRED_REPRESENTATIVE_ROLES,
    assemble_workbook,
    audit_saved_workbook,
    expected_sheet_names,
    validate_representative_assignments,
    validate_representative_payload,
    validate_role_payload_mapping,
)


SCENARIOS = [f"C{collection:03d}_G{recovery:03d}" for collection in (100, 85, 70, 50, 30) for recovery in (100, 70, 50, 30, 0)]
GLOBAL_SHEETS = [
    "README", "CONSTANTS", "SCENARIO_MATRIX", "HEATMAP", "SENSITIVITY",
    "VALIDATION_SUMMARY", "ERROR_LOG", "FINAL_CLOSURE_SUMMARY",
    "DAILY_DISTRIBUTION", "REPRESENTATIVE_CASES",
]
SUFFIXES = ["_STATS", "_DAILY", "_ASSETS", "_ACCTG", "_PARTNER", "_GUARANTEE", "_CONTRACTS"]
ROLE_MAP = {role: f"synthetic_{role}" for role in REQUIRED_REPRESENTATIVE_ROLES}


def assignments_for(scenarios=SCENARIOS, trial_id=1):
    return {
        scenario: {role: trial_id for role in REQUIRED_REPRESENTATIVE_ROLES}
        for scenario in scenarios
    }


def trial_rows_for(scenarios=SCENARIOS, trial_ids=(1, 2)):
    return {
        scenario: [{"trial_id": trial_id} for trial_id in trial_ids]
        for scenario in scenarios
    }


def payload_for(scenario, trial_id=1):
    return {
        "summary": {
            "scenario_id": scenario,
            "trial_id": trial_id,
            "final_close_date": date(2027, 1, 2),
            "final_net_project_equity": 1000,
            "final_cash": 500,
            "cumulative_project_profit": 250,
            "termination_count": 0,
            "secondary_cycle_count": 0,
            "owned_bikes": 5,
            "held_assets": 0,
            "total_operating_revenue": 2000,
            "bad_debt": 0,
            "guarantee_recovered": 0,
        },
        "daily": [{"trial_id": trial_id, "date": date(2027, 1, 1), "Cash": 500}],
    }


def test_role_payload_mapping_fails_closed_if_unresolved_or_incomplete():
    with pytest.raises(ValueError, match="STAGE13_ROLE_TO_PAYLOAD_MAPPING_UNRESOLVED"):
        validate_role_payload_mapping(None)
    with pytest.raises(ValueError, match="STAGE13_ROLE_TO_PAYLOAD_MAPPING_KEYS_INVALID"):
        validate_role_payload_mapping({role: label for role, label in ROLE_MAP.items() if role != "MaxCase"})
    duplicated = dict(ROLE_MAP)
    duplicated["MaxCase"] = duplicated["LossCase"]
    with pytest.raises(ValueError, match="STAGE13_ROLE_TO_PAYLOAD_MAPPING_LABEL_DUPLICATE"):
        validate_role_payload_mapping(duplicated)


def test_role_mapping_requires_exact_five_roles_and_preserves_explicit_labels():
    assert validate_role_payload_mapping(ROLE_MAP) == ROLE_MAP


def test_assignments_accept_exact_five_roles_and_existing_trial_ids():
    assignments = assignments_for()
    validate_representative_assignments(
        assignments,
        expected_scenarios=SCENARIOS,
        trial_rows_by_scenario=trial_rows_for(),
    )


def test_assignment_rejects_legacy_loss_alias_and_missing_maxcase():
    assignments = assignments_for()
    for scenario in SCENARIOS:
        assignments[scenario]["Loss"] = assignments[scenario].pop("LossCase")
        assignments[scenario].pop("MaxCase")
    with pytest.raises(ValueError, match="STAGE13_REPRESENTATIVE_ROLE_KEYS_INVALID"):
        validate_representative_assignments(assignments, expected_scenarios=SCENARIOS)


@pytest.mark.parametrize("bad_id", [True, 1.0, 0, -1, "1"])
def test_assignment_rejects_bool_float_or_nonpositive_trial_ids(bad_id):
    assignments = assignments_for()
    assignments[SCENARIOS[0]]["P50"] = bad_id
    with pytest.raises(ValueError, match="STAGE13_REPRESENTATIVE_TRIAL_ID_INVALID"):
        validate_representative_assignments(assignments, expected_scenarios=SCENARIOS)


def test_assignment_rejects_unknown_or_duplicated_canonical_trial_id():
    assignments = assignments_for()
    with pytest.raises(ValueError, match="STAGE13_REPRESENTATIVE_TRIAL_BINDING_INVALID"):
        validate_representative_assignments(
            assignments,
            expected_scenarios=SCENARIOS,
            trial_rows_by_scenario=trial_rows_for(trial_ids=(2, 3)),
        )
    duplicated_rows = trial_rows_for(trial_ids=(1, 1))
    with pytest.raises(ValueError, match="STAGE13_REPRESENTATIVE_TRIAL_BINDING_INVALID"):
        validate_representative_assignments(
            assignments,
            expected_scenarios=SCENARIOS,
            trial_rows_by_scenario=duplicated_rows,
        )


def test_assignment_rejects_missing_or_extra_scenario_keys():
    assignments = assignments_for()
    assignments.pop(SCENARIOS[-1])
    with pytest.raises(ValueError, match="STAGE13_REPRESENTATIVE_SCENARIO_KEYS_INVALID"):
        validate_representative_assignments(assignments, expected_scenarios=SCENARIOS)


def test_payload_binds_exact_scenario_trial_and_daily_rows():
    scenario = SCENARIOS[0]
    validate_representative_payload(scenario, "P50", 1, payload_for(scenario))


def test_payload_rejects_wrong_scenario_or_trial_id_types_and_daily_binding():
    scenario = SCENARIOS[0]
    wrong_scenario = payload_for("C030_G000")
    with pytest.raises(ValueError, match="STAGE13_REPRESENTATIVE_SCENARIO_MISMATCH"):
        validate_representative_payload(scenario, "P50", 1, wrong_scenario)
    wrong_trial = payload_for(scenario, 1)
    wrong_trial["summary"]["trial_id"] = True
    with pytest.raises(ValueError, match="STAGE13_REPRESENTATIVE_TRIAL_MISMATCH"):
        validate_representative_payload(scenario, "P50", 1, wrong_trial)
    wrong_daily = payload_for(scenario, 1)
    wrong_daily["daily"][0]["trial_id"] = 2
    with pytest.raises(ValueError, match="STAGE13_REPRESENTATIVE_DAILY_TRIAL_MISMATCH"):
        validate_representative_payload(scenario, "P50", 1, wrong_daily)


def test_expected_sheet_inventory_has_exact_185_names_in_declared_order():
    names = expected_sheet_names(GLOBAL_SHEETS, SCENARIOS, SUFFIXES)
    assert len(names) == 185
    assert names[:10] == GLOBAL_SHEETS
    assert names[10:17] == [f"{SCENARIOS[0]}{suffix}" for suffix in SUFFIXES]
    assert names[-7:] == [f"{SCENARIOS[-1]}{suffix}" for suffix in SUFFIXES]


def test_saved_xlsx_audit_reads_back_hash_size_and_rejects_formula(tmp_path: Path):
    expected = ["README", "C100_G100_STATS"]
    output = tmp_path / "audit.xlsx"
    wb = Workbook()
    wb.active.title = "README"
    wb.create_sheet("C100_G100_STATS")
    wb.save(output)
    wb.close()
    audit = audit_saved_workbook(output, expected)
    assert audit["sheet_count"] == 2
    assert audit["sheet_names_exact"] is True
    assert audit["file_bytes"] == output.stat().st_size
    assert len(audit["sha256"]) == 64

    formula_output = tmp_path / "formula.xlsx"
    wb = Workbook()
    wb.active["A1"] = "=1+1"
    wb.save(formula_output)
    wb.close()
    with pytest.raises(RuntimeError, match="STAGE13_FORMULA_CELLS"):
        audit_saved_workbook(formula_output, ["Sheet"])



def test_public_builder_fails_closed_before_parquet_read_or_xlsx_output(tmp_path: Path):
    manifest_path = tmp_path / "manifest.json"
    manifest_path.write_text(
        __import__("json").dumps({
            "source_sha": excel_builder.EXPECTED_SOURCE_SHA,
            "scenarios": {scenario: {} for scenario in SCENARIOS},
            "total_capital": 10000,
        }),
        encoding="utf-8",
    )
    assignment_path = tmp_path / "representative_trial_ids.json"
    assignment_path.write_text(
        __import__("json").dumps({
            "scenarios": {
                scenario: {"P10": 1, "P50": 1, "P90": 1, "Loss": 1}
                for scenario in SCENARIOS
            }
        }),
        encoding="utf-8",
    )
    output = tmp_path / "must-not-be-created.xlsx"
    with pytest.raises(ValueError, match="STAGE13_REPRESENTATIVE_ROLE_KEYS_INVALID"):
        excel_builder.build_workbook(
            str(tmp_path / "missing-dataset-root"),
            str(output),
            str(manifest_path),
            str(assignment_path),
        )
    assert not output.exists()


def _fake_builder(payloads):
    def make_global(name):
        def write(wb, *args, **kwargs):
            if name == "README":
                ws = wb.active
                ws.title = "README"
                ws.append(["Excel sheets", 1])
                ws.append(["Excel error count", 0])
            else:
                wb.create_sheet(name)
        return write

    def write_stats(wb, scenario, *args, **kwargs):
        wb.create_sheet(f"{scenario}_STATS")

    def write_details(wb, scenario, *args, **kwargs):
        for suffix in SUFFIXES[1:]:
            wb.create_sheet(f"{scenario}{suffix}")

    return SimpleNamespace(
        EXPECTED_SCENARIOS=SCENARIOS,
        GLOBAL_SHEETS=GLOBAL_SHEETS,
        SUFFIXES=SUFFIXES,
        write_readme=make_global("README"),
        write_constants=make_global("CONSTANTS"),
        build_scenario_summary=lambda rows: {},
        write_scenario_matrix=make_global("SCENARIO_MATRIX"),
        write_heatmap=make_global("HEATMAP"),
        write_sensitivity=make_global("SENSITIVITY"),
        write_validation=make_global("VALIDATION_SUMMARY"),
        write_errors=make_global("ERROR_LOG"),
        write_final_closure=make_global("FINAL_CLOSURE_SUMMARY"),
        write_daily_distribution=make_global("DAILY_DISTRIBUTION"),
        write_representatives=make_global("REPRESENTATIVE_CASES"),
        write_stats_sheet=write_stats,
        write_scenario_detail=write_details,
        load_rep_payload=lambda root, scenario, label: payloads[scenario]["P50"],
        save_workbook_preserving_stats_precision=lambda wb, path: (wb.save(path) or {"synthetic": True}),
    )


def test_synthetic_assembly_creates_and_reopens_exact_185_sheet_structure(tmp_path: Path):
    assignments = assignments_for()
    rows = trial_rows_for(trial_ids=(1,))
    payloads = {
        scenario: {role: payload_for(scenario, 1) for role in REQUIRED_REPRESENTATIVE_ROLES}
        for scenario in SCENARIOS
    }
    data = {
        "manifest": {"total_capital": 10000},
        "trials_by_scenario": rows,
        "representatives_by_scenario": payloads,
        "representative_trial_ids": assignments,
        "validation": {"status": "synthetic-test-only", "tests": []},
        "errors": [],
        "dataset_root": str(tmp_path),
    }
    output = tmp_path / "synthetic-assembly.xlsx"
    result = assemble_workbook(
        builder=_fake_builder(payloads),
        data=data,
        output_path=output,
        role_to_payload_labels=ROLE_MAP,
        mode="synthetic-test-only",
    )
    assert result["success"] is True
    assert result["sheets_count"] == 185
    assert result["trial_count"] == 25
    wb = load_workbook(output, data_only=False, read_only=False)
    try:
        assert wb.sheetnames == expected_sheet_names(GLOBAL_SHEETS, SCENARIOS, SUFFIXES)
        assert wb["README"]["B1"].value == 185
        assert wb["README"]["B2"].value == 0
        assert all(cell.data_type != "f" and cell.data_type != "e" for ws in wb.worksheets for row in ws.iter_rows() for cell in row)
    finally:
        wb.close()


def test_provisional_p16_uses_incremental_role_writer_and_keeps_185_sheet_order(tmp_path: Path):
    assignments = assignments_for()
    rows = trial_rows_for(trial_ids=(1,))
    payloads = {
        scenario: {role: payload_for(scenario, 1) for role in REQUIRED_REPRESENTATIVE_ROLES if role != "MaxCase"}
        for scenario in SCENARIOS
    }
    data = {
        "manifest": {"total_capital": 10000},
        "trials_by_scenario": rows,
        "representatives_by_scenario": payloads,
        "representative_trial_ids": assignments,
        "validation": {"status": "synthetic-test-only", "tests": []},
        "errors": [],
        "dataset_root": str(tmp_path),
        "pending_roles": ["MaxCase"],
    }
    fake_builder = _fake_builder(payloads)

    def start_representative_sheet(wb):
        ws = wb.create_sheet("REPRESENTATIVE_CASES")
        ws.append(["scenario", "case", "record_type", "trial_id"])
        return ws

    def append_representative_scenario(ws, scenario, role_payloads, assignment, role_map, scenario_rows, *, on_p50=None):
        ws.append([scenario, "MaxCase", "PENDING_DATA_SOURCE", assignment["MaxCase"]])
        if on_p50 is not None:
            on_p50(role_payloads["P50"])

    fake_builder.start_representative_sheet = start_representative_sheet
    fake_builder.append_representative_scenario = append_representative_scenario
    fake_builder.finish_representative_sheet = lambda ws: None

    output = tmp_path / "provisional-p16.xlsx"
    result = assemble_workbook(
        builder=fake_builder,
        data=data,
        output_path=output,
        role_to_payload_labels=ROLE_MAP,
        mode="provisional-p16",
    )
    assert result["success"] is True
    assert result["sheets_count"] == 185
    assert result["warnings"]
    workbook = load_workbook(output, data_only=True)
    try:
        assert workbook.sheetnames == expected_sheet_names(GLOBAL_SHEETS, SCENARIOS, SUFFIXES)
        representative = workbook["REPRESENTATIVE_CASES"]
        rows = list(representative.iter_rows(min_row=2, values_only=True))
        assert len(rows) == 25
        assert all(row[2] == "PENDING_DATA_SOURCE" for row in rows)
        assert any("MAXCASE_PAYLOAD_PENDING_DATA_SOURCE" in warning for warning in result["warnings"])
    finally:
        workbook.close()
