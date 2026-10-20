from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook

from scripts.excel_builder import DATE_FMT, FINANCE_FMT, write_daily_sheet


def test_daily_sheet_writes_mapping_values_in_header_order_and_round_trips(tmp_path: Path):
    payload = {
        "daily": [
            {
                "date": date(2027, 1, 1),
                "Final_Close_Date": date(2027, 1, 3),
                "Cash": 100_000,
                "Accounts_Receivable": None,
            },
            {
                "date": date(2027, 1, 2),
                "Final_Close_Date": date(2027, 1, 3),
                "Cash": 99_000,
                "Accounts_Receivable": 250,
            },
        ]
    }
    headers = list(payload["daily"][0])
    output = tmp_path / "daily.xlsx"

    wb = Workbook()
    write_daily_sheet(wb, "FIXTURE", payload)
    wb.save(output)
    wb.close()

    inspected = load_workbook(output, data_only=False)
    ws = inspected["FIXTURE_DAILY"]
    assert tuple(cell.value for cell in ws[1]) == tuple(headers)
    assert ws.max_column == len(headers)
    assert ws.max_row == len(payload["daily"]) + 1

    actual = [
        tuple(value.date() if isinstance(value, datetime) else value for value in row)
        for row in ws.iter_rows(min_row=2, values_only=True)
    ]
    expected = [tuple(row.get(header) for header in headers) for row in payload["daily"]]
    assert actual == expected
    assert type(actual[0][headers.index("Cash")]) is int
    assert actual[0][headers.index("Accounts_Receivable")] is None
    assert ws.cell(2, headers.index("date") + 1).is_date
    assert ws.cell(2, headers.index("date") + 1).number_format == DATE_FMT
    assert ws.cell(2, headers.index("Cash") + 1).number_format == FINANCE_FMT
    assert ws.freeze_panes == "B2"
    inspected.close()
