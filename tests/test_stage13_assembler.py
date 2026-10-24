from __future__ import annotations

from dataclasses import asdict
from datetime import date
from pathlib import Path
import hashlib
import json

import monte_carlo
import pytest
from openpyxl import load_workbook

from scripts.stage13_maxcase_source_audit import inventory_identity_mismatches, select_maxcase
from scripts.stage13_assemble import (
    ADAPTIVE_CHECKPOINT_METRICS, EXPECTED_ADAPTIVE_COHORT_ID,
    EXPECTED_REFERENCE_SHA256, EXPECTED_SCENARIOS, EXPECTED_SOURCE_SHA,
    EXPECTED_ADAPTIVE_SOURCE_COMMIT, _canonical_scenario_fingerprint,
    _load_canonical_dataset_identity, _metric_status,
    _validate_canonical_dataset_fingerprints, _validate_mode_and_identity,
    assemble_workbook, representative_daily_value, select_representative_trial_ids,
)

FOCUS_SCENARIOS = EXPECTED_SCENARIOS[:5]
START_DATE = date(2027, 1, 1)
END_DATE = date(2027, 1, 2)


def trial_result_sha256(row):
    trial_result = monte_carlo.TrialResult(
        **{key: value for key, value in row.items() if key != "result_sha256"}
    )
    # The Stage 13 result_sha256 covers the financial TrialResult payload;
    # RNG checksum metadata is validated separately and must not change it.
    fingerprint_payload = asdict(trial_result)
    fingerprint_payload.pop("rng_checksum", None)
    fingerprint_payload.pop("rng_draw_count", None)
    raw = json.dumps(
        fingerprint_payload,
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
        default=str,
    ).encode("utf-8")
    return hashlib.sha256(raw).hexdigest()


def trial_row(scenario, trial_id, offset):
    equity = 1000 + offset * 100 + trial_id * 10
    return {
        "scenario_id": scenario,
        "collection_probability": int(scenario[1:4]) / 100,
        "recovery_rate_pct": int(scenario[6:9]),
        "trial_id": trial_id, "final_close_date": "2027-01-02",
        "final_net_project_equity": equity, "final_cash": equity + 100,
        "cumulative_project_profit": trial_id * 100 - 250,
        "operating_net_profit": trial_id * 10, "termination_count": trial_id % 4,
        "secondary_cycle_count": trial_id % 3, "owned_bikes": 8 + trial_id % 3,
        "held_assets": trial_id % 2, "total_operating_revenue": 500 + trial_id * 20,
        "bad_debt": trial_id * 2, "guarantee_recovered": trial_id,
        "partner1_final_entitlement": equity * 7 // 10,
        "partner2_final_entitlement": equity - equity * 7 // 10,
        "rng_checksum": (
            0 if int(scenario[1:4]) / 100.0 in (0.0, 1.0)
            else (offset + 1) * 1_000_000 + trial_id
        ),
        "rng_draw_count": (
            0 if int(scenario[1:4]) / 100.0 in (0.0, 1.0) else 20
        ),
    }


def daily_rows(trial):
    equity = int(trial["final_net_project_equity"])
    return [
        {
            "date": START_DATE, "trial_id": int(trial["trial_id"]), "Final_Close_Date": END_DATE, "Cash": int(trial["final_cash"]) - 100,
            "Net_Equity": equity - 100, "Active_Bikes": 4, "Owned_Transferred_Bikes": 1,
            "Pending_Claims": 0, "Partner1_Reinvestment_Balance": 30, "Partner2_Reinvestment_Balance": 20,
            "Operating_Revenue": 250, "Operating_Expenses": 10, "Operating_Net_Profit": 240,
            "Cumulative_Project_Profit": int(trial["cumulative_project_profit"]) - 1,
            "Accounts_Receivable": 10, "Guarantee_Claim_Receivable": 0,
            "Gross_Bike_Assets": 1000, "Accumulated_Depreciation": 100,
        },
        {
            "date": END_DATE, "trial_id": int(trial["trial_id"]), "Final_Close_Date": END_DATE, "Cash": int(trial["final_cash"]),
            "Net_Equity": equity, "Active_Bikes": 0, "Owned_Transferred_Bikes": 1,
            "Pending_Claims": 0, "Partner1_Reinvestment_Balance": 0, "Partner2_Reinvestment_Balance": 0,
            "Operating_Revenue": 0, "Operating_Expenses": 0, "Operating_Net_Profit": 0,
            "Cumulative_Project_Profit": int(trial["cumulative_project_profit"]),
            "Accounts_Receivable": 0, "Guarantee_Claim_Receivable": 0,
            "Gross_Bike_Assets": 1000, "Accumulated_Depreciation": 100,
        },
    ]


def rep_payload(trial):
    return {
        "summary": dict(trial), "daily": daily_rows(trial),
        "final_inventory": [{
            "bike_id": "BIKE-001", "source": "fixture", "purchase_date": "2027-01-01",
            "delivery_date": "2027-01-01", "gross_cost": 1000, "accumulated_depreciation": 100,
            "net_book_value": 900, "termination_count": trial["termination_count"],
            "secondary_cycle_count": trial["secondary_cycle_count"], "state": "CLOSED",
        }],
        "accounting": {"daily_balance_checks": [], "cash_rollforward": [], "ar_rollforward": [],
                       "asset_rollforward": [], "equity_rollforward": []},
        "partner_memo": [
            {"date": START_DATE, "Partner1_Reinvestment_Balance": 30, "Partner2_Reinvestment_Balance": 20},
            {"date": END_DATE, "Partner1_Reinvestment_Balance": 0, "Partner2_Reinvestment_Balance": 0},
        ],
        "final_reconciliation": {}, "guarantee_claims": [], "contracts": [],
    }


def demo_data():
    trials = {scenario: [] for scenario in EXPECTED_SCENARIOS}
    distributions = {scenario: [] for scenario in EXPECTED_SCENARIOS}
    representatives, assignments = {}, {}
    total = 0
    for offset, scenario in enumerate(FOCUS_SCENARIOS):
        rows = [trial_row(scenario, trial_id, offset) for trial_id in range(1, 11)]
        trials[scenario] = rows
        total += len(rows)
        selected = select_representative_trial_ids(rows, scenario)
        assignments[scenario] = selected
        by_id = {row["trial_id"]: row for row in rows}
        representatives[scenario] = {
            role: rep_payload(by_id[trial_id]) for role, trial_id in selected.items()
        }
        distributions[scenario] = [{
            "date": START_DATE, "active_trial_count": 10,
            "P10_Cash": 100, "P25_Cash": 110, "P50_Cash": 120, "P75_Cash": 130, "P90_Cash": 140,
            "P10_Net_Equity": 90, "P25_Net_Equity": 100, "P50_Net_Equity": 110, "P75_Net_Equity": 120, "P90_Net_Equity": 130,
            "P10_Active_Bikes": 1, "P25_Active_Bikes": 2, "P50_Active_Bikes": 3, "P75_Active_Bikes": 4, "P90_Active_Bikes": 5,
            "P10_Owned_Transferred_Bikes": 0, "P25_Owned_Transferred_Bikes": 0,
            "P50_Owned_Transferred_Bikes": 1, "P75_Owned_Transferred_Bikes": 1, "P90_Owned_Transferred_Bikes": 2,
            "P10_Pending_Claims": 0, "P25_Pending_Claims": 0, "P50_Pending_Claims": 0,
            "P75_Pending_Claims": 1, "P90_Pending_Claims": 1,
        }]
    return {
        "manifest": {
            "source_sha": EXPECTED_SOURCE_SHA, "run_id": "TEST-DEMO",
            "run_head_sha": "synthetic-test-only", "master_seed": 20270101,
            "total_capital": 10000, "total_trials": total, "total_representatives": 25,
            "global_manifest_sha256": "synthetic-fixture-not-production-evidence",
            "constant_map": {"MASTER_SEED": 20270101, "TOTAL_CAPITAL": 10000},
            "scenarios": {s: {"final_status": "DEMO_ONLY"} for s in EXPECTED_SCENARIOS},
        },
        "trials_by_scenario": trials, "daily_distributions_by_scenario": distributions,
        "representatives_by_scenario": representatives, "representative_trial_ids": assignments,
        "validation": {"status": "DEMO_ONLY", "tests": []}, "errors": [],
    }


def synthetic_canonical_fingerprint_data():
    rows_by_scenario = {}
    per_scenario_identity = {}
    for offset, scenario in enumerate(EXPECTED_SCENARIOS):
        max_trial_id = 1 if scenario.startswith("C100_") else 550
        rows = [
            trial_row(scenario, trial_id, offset)
            for trial_id in range(1, max_trial_id + 1)
        ]
        rows_by_scenario[scenario] = rows
        fingerprint = _canonical_scenario_fingerprint(rows)
        per_scenario_identity[scenario] = {
            "scenario_fingerprint_stage9a": fingerprint,
            "combined_dataset_fingerprint_sha256_stage10": (
                None if scenario.startswith("C100_") else fingerprint
            ),
        }
    return rows_by_scenario, {"per_scenario": per_scenario_identity}


def test_canonical_identity_loader_requires_the_reviewed_stage11_blob():
    identity = _load_canonical_dataset_identity()
    assert identity["baseline_commit"] == "17f4e8d0ddc1908e2c34bf927ef1d1b7a1c4bf30"
    assert identity["total_canonical_trials"] == 11005
    assert identity["scenarios_count"] == 25
    assert set(identity["per_scenario"]) == set(EXPECTED_SCENARIOS)
    assert identity["sources"]["stage9a"]["global_manifest_sha256_9005"] == (
        "892da438955160bb59aff3046271402f57bdd61a4abb35f41db7188f61346ff6"
    )


def test_canonical_identity_loader_rejects_modified_identity_blob(tmp_path: Path):
    original = (
        Path(__file__).resolve().parents[1]
        / "docs"
        / "stage11"
        / "dataset_identity.json"
    ).read_bytes()
    altered = tmp_path / "dataset_identity.json"
    altered.write_bytes(original + b" ")

    with pytest.raises(
        RuntimeError,
        match="CANONICAL_DATASET_IDENTITY_BLOB_MISMATCH",
    ):
        _load_canonical_dataset_identity(altered)


def test_canonical_fingerprint_uses_the_existing_stage7_stage10_float_encoding():
    assert _canonical_scenario_fingerprint([{"trial_id": 1, "value": 1.5}]) == (
        "823067f9ffc1d4067bc35a6da17521881a62fca027f80090f922cd9d9f835425"
    )
    # Multi-row fixture independently pins both the float.hex() encoding and
    # the single literal newline separator between canonical trial rows.
    assert _canonical_scenario_fingerprint([
        {"trial_id": 2, "value": 2.0},
        {"trial_id": 1, "value": 1.5},
    ]) == (
        "777a0cd93f188002deb11adda93e1ceb2d297235ea4548c1faeeefeabf2836f8"
    )


def test_canonical_per_scenario_fingerprints_accept_matching_rows():
    rows_by_scenario, identity = synthetic_canonical_fingerprint_data()
    _validate_canonical_dataset_fingerprints(rows_by_scenario, identity)


def test_canonical_per_scenario_fingerprints_reject_mutated_financial_value():
    rows_by_scenario, identity = synthetic_canonical_fingerprint_data()
    # Preserve scenario/Trial ID and the identity manifest; only result data changes.
    rows_by_scenario["C085_G100"][0]["final_cash"] += 1
    with pytest.raises(
        RuntimeError,
        match="CANONICAL_SCENARIO_FINGERPRINT_MISMATCH: C085_G100",
    ):
        _validate_canonical_dataset_fingerprints(rows_by_scenario, identity)


def test_canonical_mode_integrates_pinned_identity_and_fails_on_value_drift(monkeypatch):
    import scripts.stage13_assemble as assembler_module

    rows_by_scenario, identity = synthetic_canonical_fingerprint_data()
    monkeypatch.setattr(
        assembler_module,
        "_load_canonical_dataset_identity",
        lambda: identity,
    )
    data = {
        "manifest": {
            "source_sha": EXPECTED_SOURCE_SHA,
            "master_seed": 20270101,
            "total_trials": 11005,
            "scenarios": {scenario: {} for scenario in EXPECTED_SCENARIOS},
        },
        "trials_by_scenario": rows_by_scenario,
    }

    _validate_mode_and_identity(data, mode="canonical")
    data["trials_by_scenario"]["C085_G100"][0]["final_cash"] += 1
    with pytest.raises(
        RuntimeError,
        match="CANONICAL_SCENARIO_FINGERPRINT_MISMATCH: C085_G100",
    ):
        _validate_mode_and_identity(data, mode="canonical")


def test_assemble_185_sheets_and_five_representative_roles(tmp_path: Path):
    data = demo_data()
    output = tmp_path / "demo.xlsx"
    result = assemble_workbook(data, output, mode="demo")
    assert result["success"] is True
    assert result["trial_count"] == 50
    assert (result["sheets_count"], result["stats_sheets"], result["daily_sheets"]) == (185, 25, 25)
    assert result["excel_errors"] == 0
    assert result["formula_cells"] == 0
    wb = load_workbook(output, data_only=False)
    assert len(wb.sheetnames) == 185
    assert wb["C100_G100_STATS"]["A2"].value == "Final_Net_Project_Equity"
    assert wb["C100_G100_DAILY"].max_row == 3
    rep = wb["REPRESENTATIVE_CASES"]
    summaries = [tuple(row) for row in rep.iter_rows(min_row=2, values_only=True)
                 if row[0] in FOCUS_SCENARIOS and row[2] == "SUMMARY"]
    assert len(summaries) == 25
    assert {row[1] for row in summaries} == {"P10", "P50", "P90", "Loss Case", "Max Case"}
    assert all(cell.data_type not in {"f", "e"} for ws in wb.worksheets
               for row in ws.iter_rows() for cell in row)
    wb.close()


def test_loss_and_max_case_selection_matches_normative_decision():
    rows = [trial_row("C100_G100", i, 0) for i in range(1, 11)]
    selected = select_representative_trial_ids(rows, "C100_G100")
    assert selected["Loss"] == 1
    assert selected["Max Case"] == 10


def test_max_case_without_positive_equity_is_undefined_business_rule():
    rows = [trial_row("C100_G100", i, 0) for i in range(1, 4)]
    for row in rows:
        row["final_net_project_equity"] = -abs(row["final_net_project_equity"])
    with pytest.raises(RuntimeError, match="UNDEFINED_BUSINESS_RULE: MaxCase has no positive"):
        select_representative_trial_ids(rows, "C100_G100")


def test_incorrect_recorded_representative_trial_id_is_rejected():
    data = demo_data()
    data["representative_trial_ids"]["C100_G100"]["P50"] = 999
    with pytest.raises(RuntimeError, match="REPRESENTATIVE_TRIAL_ID_MISMATCH"):
        assemble_workbook(data, Path("unused.xlsx"), mode="demo")


def test_canonical_mode_rejects_demo_subset(tmp_path: Path):
    data = demo_data()
    output = tmp_path / "must-not-exist.xlsx"
    with pytest.raises(RuntimeError):
        assemble_workbook(data, output, mode="canonical")
    assert not output.exists()

def test_sensitivity_sheet_holds_other_dimension_fixed_and_daily_stats_are_values(tmp_path: Path):
    data = demo_data()
    output = tmp_path / "sensitivity.xlsx"
    result = assemble_workbook(data, output, mode="demo")
    assert result["success"] is True
    assert result["sheets_count"] == 185

    wb = load_workbook(output, data_only=True)
    assert wb.sheetnames.count("SENSITIVITY") == 1
    ws = wb["SENSITIVITY"]
    assert [cell.value for cell in ws[1]] == [
        "varied_dimension", "fixed_dimension", "fixed_value_pct",
        "varied_value_pct", "scenario", "mean_cumulative_project_profit",
    ]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert len(rows) == 50
    collection_rows = [row for row in rows if row[0] == "collection_probability_pct"]
    recovery_rows = [row for row in rows if row[0] == "guarantee_recovery_rate_pct"]
    assert len(collection_rows) == 25
    assert len(recovery_rows) == 25
    for row in collection_rows:
        collection_pct = int(row[4].split("_")[0][1:])
        recovery_pct = int(row[4].split("_")[1][1:])
        assert row[1] == "guarantee_recovery_rate_pct"
        assert row[2] == recovery_pct
        assert row[3] == collection_pct
    for row in recovery_rows:
        collection_pct = int(row[4].split("_")[0][1:])
        recovery_pct = int(row[4].split("_")[1][1:])
        assert row[1] == "collection_probability_pct"
        assert row[2] == collection_pct
        assert row[3] == recovery_pct

    target = [row for row in rows if row[4] == "C100_G100"]
    assert len(target) == 2
    assert all(row[5] == 300 for row in target)

    stats = wb["C100_G100_STATS"]
    assert stats["A2"].value == "Final_Net_Project_Equity"
    assert stats["B2"].value == 1055.0
    daily = wb["C100_G100_DAILY"]
    assert daily["A2"].value.date() == START_DATE
    assert daily["B2"].value.date() == END_DATE
    assert daily["C2"].value == 1050
    wb.close()


def adaptive_data(stable_counts=None):
    """Build synthetic adaptive-cohort evidence; does not run the simulation engine."""
    base = demo_data()
    stable_counts = dict(stable_counts or {})
    for scenario in EXPECTED_SCENARIOS:
        stable_counts.setdefault(scenario, 600)

    trials = {}
    distributions = {}
    representatives = {}
    assignments = {}
    scenario_evidence = {}
    total_trials = 0

    for offset, scenario in enumerate(EXPECTED_SCENARIOS):
        stable_at_n = stable_counts[scenario]
        rows = [trial_row(scenario, trial_id, offset) for trial_id in range(1, stable_at_n + 1)]
        for row in rows:
            row["result_sha256"] = trial_result_sha256(row)
        trials[scenario] = rows
        total_trials += len(rows)

        selected = select_representative_trial_ids(rows, scenario)
        assignments[scenario] = selected
        by_id = {row["trial_id"]: row for row in rows}
        representatives[scenario] = {
            role: rep_payload(by_id[trial_id]) for role, trial_id in selected.items()
        }
        distributions[scenario] = [dict(base["daily_distributions_by_scenario"][FOCUS_SCENARIOS[0]][0])]

        history = [{
            "n": 300,
            "checkpoint_status": "COVERAGE_ONLY",
            "overall_status": "COVERAGE_ONLY",
            "section_2_5_status": "COVERAGE_ONLY",
        }]
        if not scenario.startswith("C100_"):
            history[0]["p50_final_net_project_equity"] = 1000.0
        n = 600
        while n <= stable_at_n:
            metric_results = {}
            should_fail = scenario == "C030_G030" and n == 600 and stable_at_n > 600
            for metric in ADAPTIVE_CHECKPOINT_METRICS:
                old_value = 0.2 if metric == "Probability_of_Accounting_Loss" else 1000.0
                if should_fail and metric == "P90_Final_Cash":
                    new_value = old_value * 1.02
                elif metric == "Probability_of_Accounting_Loss":
                    new_value = old_value * 1.002
                else:
                    new_value = old_value * 1.002
                metric_results[metric] = {
                    "old_value": old_value,
                    "new_value": new_value,
                    "status": _metric_status(old_value, new_value),
                }
            section_status = (
                "PASS" if all(x["status"] == "PASS" for x in metric_results.values()) else "FAIL"
            )
            entry = {
                "n": n,
                "previous_n": n // 2,
                "overall_status": "SCENARIO_STABLE" if n == stable_at_n else "NOT_STABLE",
                "gate_a_status": "PASS",
                "gate_a_evidence_ref": f"synthetic-gate-a:{scenario}:{n}",
                "gate_a_evidence_sha256": hashlib.sha256(f"gate-a-evidence|{scenario}|{n}".encode()).hexdigest(),
                "gate_b_status": "PASS",
                "gate_b_evidence_ref": f"synthetic-gate-b:{scenario}:{n}",
                "gate_b_evidence_sha256": hashlib.sha256(f"gate-b-evidence|{scenario}|{n}".encode()).hexdigest(),
                "section_2_5_status": section_status,
                "section_2_5_evidence_ref": f"synthetic-section-2-5:{scenario}:{n}",
                "section_2_5_evidence_sha256": hashlib.sha256(f"section-2-5-evidence|{scenario}|{n}".encode()).hexdigest(),
                "section_2_5_metrics": metric_results,
            }
            if not scenario.startswith("C100_"):
                old_p50 = 1000.0 if n == 600 else 1010.0
                new_p50 = old_p50 * 1.01
                entry["gate_b_p50_previous"] = old_p50
                entry["gate_b_p50_current"] = new_p50
                entry["p50_final_net_project_equity"] = new_p50
                entry["gate_b_p50_continuation_required"] = False
            history.append(entry)
            n *= 2

        final_entry = history[-1]
        scenario_evidence[scenario] = {
            "final_status": "SCENARIO_STABLE",
            "stable_at_n": stable_at_n,
            "gate_a_status": "PASS",
            "gate_a_evidence_ref": f"synthetic-gate-a:{scenario}:{stable_at_n}",
            "gate_a_evidence_sha256": hashlib.sha256(f"gate-a-evidence|{scenario}|{stable_at_n}".encode()).hexdigest(),
            "gate_b_status": "PASS",
            "gate_b_evidence_ref": f"synthetic-gate-b:{scenario}:{stable_at_n}",
            "gate_b_evidence_sha256": hashlib.sha256(f"gate-b-evidence|{scenario}|{stable_at_n}".encode()).hexdigest(),
            "section_2_5_status": "PASS",
            "section_2_5_evidence_ref": f"synthetic-section-2-5:{scenario}:{stable_at_n}",
            "section_2_5_evidence_sha256": hashlib.sha256(f"section-2-5-evidence|{scenario}|{stable_at_n}".encode()).hexdigest(),
            "checkpoint_history": history,
        }
        if scenario.startswith("C100_"):
            scenario_evidence[scenario]["c100_deterministic_equivalence_status"] = "PASS"
            scenario_evidence[scenario]["c100_deterministic_evidence_ref"] = (
                f"synthetic-c100-equivalence:{scenario}"
            )
            scenario_evidence[scenario]["c100_deterministic_evidence_sha256"] = (
                hashlib.sha256(f"c100-equivalence-evidence|{scenario}".encode()).hexdigest()
            )

    ordered_rows = [
        f"{scenario}|{row['trial_id']}|{row['result_sha256']}\n"
        for scenario in EXPECTED_SCENARIOS
        for row in trials[scenario]
    ]
    ordered_sha = hashlib.sha256("".join(ordered_rows).encode("utf-8")).hexdigest()
    ordered_rng_audit_rows = [
        f"{scenario}|{row['trial_id']}|{row['rng_checksum']}|{row['rng_draw_count']}\n"
        for scenario in EXPECTED_SCENARIOS
        for row in trials[scenario]
    ]
    rng_audit_sha = hashlib.sha256(
        "".join(ordered_rng_audit_rows).encode("utf-8")
    ).hexdigest()
    manifest = dict(base["manifest"])
    manifest.update({
        "cohort_id": EXPECTED_ADAPTIVE_COHORT_ID,
        "source_sha": EXPECTED_SOURCE_SHA,
        "source_commit": EXPECTED_ADAPTIVE_SOURCE_COMMIT,
        "master_seed": 20270101,
        "reference_sha256": EXPECTED_REFERENCE_SHA256,
        "schema_version": "stage12-daily-accounting-v1",
        "rng_identity": "rng.py@e50d0410729aeeb657d4b187c0bc32f9527442d1",
        "sampling_policy": "per-scenario-doubling-300-600-1200-first-passing",
        "run_id": 99999999,
        "run_head_sha": "b" * 40,
        "total_trials": total_trials,
        "total_representatives": len(EXPECTED_SCENARIOS) * 5,
        "ordered_trial_hash_map_sha256": ordered_sha,
        "gate_a_ordered_trial_hash_map_sha256": ordered_sha,
        "rng_audit_map_sha256": rng_audit_sha,
        "scenarios": scenario_evidence,
        "dataset_status": "SYNTHETIC_TEST_ONLY",
    })
    return {
        "manifest": manifest,
        "trials_by_scenario": trials,
        "daily_distributions_by_scenario": distributions,
        "representatives_by_scenario": representatives,
        "representative_trial_ids": assignments,
        "validation": {"status": "SYNTHETIC_TEST_ONLY", "tests": []},
        "errors": [],
    }


def test_adaptive_metric_thresholds_are_strict_and_near_zero_uses_absolute_unit():
    assert _metric_status(100.0, 100.999) == "PASS"
    assert _metric_status(100.0, 101.0) == "FAIL"
    assert _metric_status(0.0, 0.999) == "PASS"
    assert _metric_status(0.0, 1.0) == "FAIL"
    with pytest.raises(RuntimeError, match="ADAPTIVE_METRIC_VALUE_MISSING_OR_NONFINITE"):
        _metric_status(float("nan"), 0.0)


def test_adaptive_validator_accepts_heterogeneous_first_passing_stops():
    data = adaptive_data({"C030_G030": 1200})
    rows = _validate_mode_and_identity(data, mode="adaptive")
    assert len(rows) == 25 * 600 + 600
    assert data["manifest"]["scenarios"]["C030_G030"]["stable_at_n"] == 1200
    assert data["manifest"]["scenarios"]["C100_G100"]["checkpoint_history"][0]["overall_status"] == "COVERAGE_ONLY"


def test_adaptive_mode_rejects_stable_at_300():
    data = adaptive_data()
    data["manifest"]["scenarios"]["C100_G100"]["stable_at_n"] = 300
    with pytest.raises(RuntimeError, match="ADAPTIVE_STABLE_AT_N_INVALID"):
        _validate_mode_and_identity(data, mode="adaptive")


def test_adaptive_mode_rejects_missing_gate_evidence():
    data = adaptive_data()
    del data["manifest"]["scenarios"]["C100_G100"]["gate_b_status"]
    with pytest.raises(RuntimeError, match="ADAPTIVE_FINAL_STATUS_FIELD_MISMATCH"):
        _validate_mode_and_identity(data, mode="adaptive")


@pytest.mark.parametrize("gate_a_status", ["BLOCKED", "FAIL", "PENDING"])
def test_adaptive_mode_rejects_nonpass_historical_gate_a_before_later_pass(
    gate_a_status,
):
    data = adaptive_data({"C085_G100": 1200})
    history = data["manifest"]["scenarios"]["C085_G100"]["checkpoint_history"]
    history[1]["gate_a_status"] = gate_a_status
    # The later N=1200 checkpoint still claims PASS and SCENARIO_STABLE.
    # A previous non-PASS Gate A must terminally block this assembled cohort.
    with pytest.raises(
        RuntimeError,
        match="ADAPTIVE_GATE_A_CHECKPOINT_NOT_PASS: C085_G100:600:",
    ):
        _validate_mode_and_identity(data, mode="adaptive")


def test_adaptive_mode_rejects_checkpoint_gap():
    data = adaptive_data({"C030_G030": 1200})
    data["manifest"]["scenarios"]["C030_G030"]["checkpoint_history"].pop(1)
    with pytest.raises(RuntimeError, match="ADAPTIVE_CHECKPOINT_HISTORY_LENGTH_MISMATCH"):
        _validate_mode_and_identity(data, mode="adaptive")


def test_adaptive_mode_rejects_nonfinite_metric_value():
    data = adaptive_data()
    final_entry = data["manifest"]["scenarios"]["C100_G100"]["checkpoint_history"][-1]
    final_entry["section_2_5_metrics"]["P90_Final_Cash"]["new_value"] = float("nan")
    with pytest.raises(RuntimeError, match="ADAPTIVE_METRIC_VALUE_MISSING_OR_NONFINITE"):
        _validate_mode_and_identity(data, mode="adaptive")


def test_adaptive_mode_recomputes_aggregate_trial_fingerprint_map():
    data = adaptive_data()
    data["manifest"]["ordered_trial_hash_map_sha256"] = "a" * 64
    with pytest.raises(RuntimeError, match="ADAPTIVE_ORDERED_MAP_SHA_MISMATCH"):
        _validate_mode_and_identity(data, mode="adaptive")


def test_adaptive_mode_rejects_trial_result_value_with_detached_sha256():
    data = adaptive_data()
    row = data["trials_by_scenario"]["C030_G030"][0]
    row["final_net_project_equity"] += 1
    with pytest.raises(RuntimeError, match="ADAPTIVE_TRIAL_FINGERPRINT_MISMATCH"):
        _validate_mode_and_identity(data, mode="adaptive")


def test_adaptive_mode_rejects_trial_result_schema_drift():
    data = adaptive_data()
    row = data["trials_by_scenario"]["C030_G030"][0]
    del row["collection_probability"]
    with pytest.raises(RuntimeError, match="ADAPTIVE_TRIAL_RESULT_SCHEMA_MISMATCH"):
        _validate_mode_and_identity(data, mode="adaptive")


@pytest.mark.parametrize(
    ("field", "wrong_value"),
    [
        ("collection_probability", 0.70),
        ("recovery_rate_pct", 70),
    ],
)
def test_adaptive_mode_rejects_scenario_parameter_mismatch_with_recomputed_hashes(
    field, wrong_value
):
    data = adaptive_data()
    row = data["trials_by_scenario"]["C085_G100"][0]
    row[field] = wrong_value
    row["result_sha256"] = trial_result_sha256(row)

    # Make the row digest, ordered hash map, and Gate A map claim consistent
    # with the altered row. Validation must still reject the scenario binding.
    ordered_rows = [
        f"{scenario}|{trial['trial_id']}|{trial['result_sha256']}\n"
        for scenario in EXPECTED_SCENARIOS
        for trial in data["trials_by_scenario"][scenario]
    ]
    ordered_sha = hashlib.sha256("".join(ordered_rows).encode("utf-8")).hexdigest()
    data["manifest"]["ordered_trial_hash_map_sha256"] = ordered_sha
    data["manifest"]["gate_a_ordered_trial_hash_map_sha256"] = ordered_sha

    with pytest.raises(
        RuntimeError,
        match="ADAPTIVE_TRIAL_SCENARIO_PARAMETER_MISMATCH",
    ):
        _validate_mode_and_identity(data, mode="adaptive")


def test_adaptive_mode_rejects_incomplete_trial_id_range():
    data = adaptive_data()
    data["trials_by_scenario"]["C100_G100"][10]["trial_id"] = 9999
    with pytest.raises(RuntimeError, match="ADAPTIVE_TRIAL_COVERAGE_MISMATCH"):
        _validate_mode_and_identity(data, mode="adaptive")


def test_adaptive_mode_enforces_gate_b_p50_2400_continuation():
    data = adaptive_data({"C030_G000": 1200})
    history = data["manifest"]["scenarios"]["C030_G000"]["checkpoint_history"]
    entry_600 = history[1]
    entry_600["gate_b_p50_current"] = 1060.0
    entry_600["p50_final_net_project_equity"] = 1060.0
    entry_600["gate_b_p50_continuation_required"] = True

    entry_1200 = history[2]
    entry_1200["gate_b_p50_previous"] = 1060.0
    entry_1200["gate_b_p50_current"] = 1070.6
    entry_1200["p50_final_net_project_equity"] = 1070.6
    entry_1200["gate_b_p50_continuation_required"] = False

    with pytest.raises(RuntimeError, match="ADAPTIVE_GATE_B_P50_REQUIRES_N2400"):
        _validate_mode_and_identity(data, mode="adaptive")


def test_adaptive_assembler_writes_185_sheets_and_truthful_readme(tmp_path: Path):
    data = adaptive_data({"C030_G030": 1200})
    output = tmp_path / "adaptive.xlsx"
    result = assemble_workbook(data, output, mode="adaptive")
    assert result["success"] is True
    assert result["trial_count"] == 25 * 600 + 600
    assert result["sheets_count"] == 185
    workbook = load_workbook(output, data_only=False)
    assert len(workbook.sheetnames) == 185
    readme = {row[0]: row[1] for row in workbook["README"].iter_rows(min_row=2, max_col=2, values_only=True)}
    assert readme["Excel sheets"] == 185
    assert readme["Excel error cells"] == 0
    assert "Stage 13 adaptive cohort" in readme["Workbook build provenance"]
    scenario_rows = {
        row[0]: row for row in workbook["SCENARIO_MATRIX"].iter_rows(min_row=2, values_only=True)
    }
    assert scenario_rows["C030_G030"][1] == 1200
    assert scenario_rows["C030_G030"][2] == "SCENARIO_STABLE"
    assert scenario_rows["C100_G100"][1] == 600
    assert scenario_rows["C100_G100"][2] == "SCENARIO_STABLE"
    assert all(
        cell.data_type not in {"f", "e"}
        for sheet in workbook.worksheets
        for row in sheet.iter_rows()
        for cell in row
    )
    workbook.close()


def test_missing_final_status_does_not_default_to_stable(tmp_path: Path):
    data = demo_data()
    del data["manifest"]["scenarios"]["C100_G100"]["final_status"]
    output = tmp_path / "status-not-inferred.xlsx"
    result = assemble_workbook(data, output, mode="demo")
    assert result["success"] is True
    workbook = load_workbook(output, data_only=True)
    rows = list(workbook["SCENARIO_MATRIX"].iter_rows(min_row=2, values_only=True))
    status_by_scenario = {row[0]: row[2] for row in rows}
    assert status_by_scenario["C100_G100"] == "NOT_ESTABLISHED"
    workbook.close()


def test_adaptive_mode_requires_gate_a_digest_to_match_the_computed_cohort_map():
    data = adaptive_data()
    data["manifest"]["gate_a_ordered_trial_hash_map_sha256"] = "a" * 64
    with pytest.raises(RuntimeError, match="ADAPTIVE_GATE_A_MAP_IDENTITY_MISMATCH"):
        _validate_mode_and_identity(data, mode="adaptive")


def test_adaptive_mode_rejects_unpinned_gate_evidence_digest():
    data = adaptive_data()
    data["manifest"]["scenarios"]["C100_G100"]["gate_a_evidence_sha256"] = "not-a-sha256"
    with pytest.raises(RuntimeError, match="ADAPTIVE_FINAL_EVIDENCE_SHA_INVALID"):
        _validate_mode_and_identity(data, mode="adaptive")


def test_adaptive_mode_rejects_p50_value_not_linked_to_previous_checkpoint():
    data = adaptive_data()
    entry = data["manifest"]["scenarios"]["C030_G000"]["checkpoint_history"][1]
    entry["gate_b_p50_previous"] = 999.0
    with pytest.raises(RuntimeError, match="ADAPTIVE_GATE_B_P50_VALUE_LINK_MISMATCH"):
        _validate_mode_and_identity(data, mode="adaptive")

def test_source_audit_maxcase_selects_highest_positive_equity():
    rows = [
        {"trial_id": 1, "final_net_project_equity": 100},
        {"trial_id": 2, "final_net_project_equity": 900},
        {"trial_id": 3, "final_net_project_equity": 500},
    ]
    assert select_maxcase(rows) == rows[1]


def test_source_audit_maxcase_tie_breaks_on_smallest_trial_id():
    rows = [
        {"trial_id": 17, "final_net_project_equity": 900},
        {"trial_id": 4, "final_net_project_equity": 900},
        {"trial_id": 2, "final_net_project_equity": 899},
    ]
    assert select_maxcase(rows) == rows[1]


def test_source_audit_maxcase_fails_closed_without_positive_integer_equity():
    rows = [
        {"trial_id": 1, "final_net_project_equity": 0},
        {"trial_id": 2, "final_net_project_equity": -1},
        {"trial_id": 3, "final_net_project_equity": 10.5},
    ]
    assert select_maxcase(rows) is None

def test_source_audit_rejects_representative_inventory_artifact_id_drift():
    # These IDs are pinned from the independently inspected GitHub artifact records:
    # V2 = artifact 11503191706 / run 37660610243; V3 = artifact 11503510055 / run 37662891608.
    manifest = {
        "artifact_manifest": {
            "enriched_artifacts": {"C050_G000": 11503191706},
            "enrichment_run_ids": {"C050_G000": 37660610243},
        }
    }
    inventory = {
        "artifacts": [{
            "scenario": "C050_G000",
            "artifact_id": 11503510055,
            "run_id": 37662891608,
        }]
    }
    assert inventory_identity_mismatches(manifest, inventory) == [{
        "scenario": "C050_G000",
        "manifest_artifact_id": 11503191706,
        "inventory_artifact_id": 11503510055,
        "manifest_run_id": 37660610243,
        "inventory_run_id": 37662891608,
    }]


def test_source_audit_rejects_v2_artifact_paired_with_v3_run():
    # Main's historical pair is internally inconsistent: the V2 artifact is owned by the V2 run.
    manifest = {
        "artifact_manifest": {
            "enriched_artifacts": {"C050_G000": 11503191706},
            "enrichment_run_ids": {"C050_G000": 37662891608},
        }
    }
    verified_v2_inventory = {
        "artifacts": [{
            "scenario": "C050_G000",
            "artifact_id": 11503191706,
            "run_id": 37660610243,
        }]
    }
    assert inventory_identity_mismatches(manifest, verified_v2_inventory) == [{
        "scenario": "C050_G000",
        "manifest_artifact_id": 11503191706,
        "inventory_artifact_id": 11503191706,
        "manifest_run_id": 37662891608,
        "inventory_run_id": 37660610243,
    }]


def test_source_audit_accepts_verified_v2_identity():
    # Exact V2 artifact/run pair confirmed in GitHub artifact metadata.
    manifest = {
        "artifact_manifest": {
            "enriched_artifacts": {"C050_G000": 11503191706},
            "enrichment_run_ids": {"C050_G000": 37660610243},
        }
    }
    verified_v2_inventory = {
        "artifacts": [{
            "scenario": "C050_G000",
            "artifact_id": 11503191706,
            "run_id": 37660610243,
        }]
    }
    assert (verified_v2_inventory["artifacts"][0]["artifact_id"],
            verified_v2_inventory["artifacts"][0]["run_id"]) == (11503191706, 37660610243)
    assert inventory_identity_mismatches(manifest, verified_v2_inventory) == []


def test_source_audit_accepts_verified_v3_identity():
    # Exact V3 artifact/run pair confirmed in GitHub artifact metadata.
    manifest = {
        "artifact_manifest": {
            "enriched_artifacts": {"C050_G000": 11503510055},
            "enrichment_run_ids": {"C050_G000": 37662891608},
        }
    }
    verified_v3_inventory = {
        "artifacts": [{
            "scenario": "C050_G000",
            "artifact_id": 11503510055,
            "run_id": 37662891608,
        }]
    }
    assert (verified_v3_inventory["artifacts"][0]["artifact_id"],
            verified_v3_inventory["artifacts"][0]["run_id"]) == (11503510055, 37662891608)
    assert inventory_identity_mismatches(manifest, verified_v3_inventory) == []

def test_maxcase_cumulative_project_profit_uses_normative_retained_earnings_fallback():
    daily = {"Retained_Earnings": -12500}
    assert representative_daily_value("Max Case", daily, "Cumulative_Project_Profit") == -12500


def test_maxcase_explicit_cumulative_project_profit_takes_precedence():
    daily = {"Cumulative_Project_Profit": 1234, "Retained_Earnings": 5678}
    assert representative_daily_value("Max Case", daily, "Cumulative_Project_Profit") == 1234


def test_maxcase_partner_memo_balances_remain_pending_when_source_is_absent():
    daily = {"Retained_Earnings": 1234}
    assert representative_daily_value("Max Case", daily, "Partner1_Reinvestment_Balance") == "PENDING_DATA_SOURCE"
    assert representative_daily_value("Max Case", daily, "Partner2_Reinvestment_Balance") == "PENDING_DATA_SOURCE"



def test_diagnostic_scenario_detail_marks_missing_source_sections_without_fabricating_values():
    from openpyxl import Workbook
    from scripts.stage13_assemble import write_scenario_detail

    wb = Workbook()
    payload = {
        "summary": {"trial_id": 7},
        "daily": [
            {
                "date": "2027-01-01",
                "trial_id": 7,
                "Cash": 120,
                "Net_Equity": 135,
                "Cumulative_Project_Profit": 15,
                "Active_Bikes": 1,
                "Owned_Transferred_Bikes": 0,
                "Pending_Claims": 0,
            }
        ],
    }
    write_scenario_detail(wb, "C085_G070", payload, diagnostic_only=True)
    assert set(wb.sheetnames) == {
        "Sheet", "C085_G070_DAILY", "C085_G070_ASSETS", "C085_G070_ACCTG",
        "C085_G070_PARTNER", "C085_G070_GUARANTEE", "C085_G070_CONTRACTS",
    }
    daily = wb["C085_G070_DAILY"]
    assert daily["A2"].value == date(2027, 1, 1)
    for suffix in ("_ASSETS", "_ACCTG", "_PARTNER", "_GUARANTEE", "_CONTRACTS"):
        ws = wb[f"C085_G070{suffix}"]
        assert ws["A2"].value == "SOURCE_NOT_AVAILABLE_IN_DIAGNOSTIC_COHORT"
        assert "no values were inferred or zero-filled" in ws["C2"].value
    wb.close()


def test_diagnostic_readme_discloses_multi_run_cohort_and_does_not_claim_stage9a_run():
    from openpyxl import Workbook
    from scripts.stage13_assemble import write_readme

    wb = Workbook()
    manifest = {
        "diagnostic_only": True,
        "cohort_id": "STAGE13_DIAGNOSTIC_CANONICAL_11005_V1",
        "total_trials": 11005,
        "total_representatives": 125,
        "master_seed": 20270101,
        "source_sha": "a" * 40,
        "run_id": 37671448546,
        "run_head_sha": "b" * 40,
        "reference_sha256": "c" * 64,
        "global_manifest_sha256": "d" * 64,
        "dataset_status": "STABLE",
        "source_artifacts": {
            "stage7_run_id": 37225915321,
            "stage7_head_sha": "e" * 40,
            "stage10_run_id": 37709936740,
            "stage10_head_sha": "f" * 40,
            "stage12_run_id": 37964656154,
            "stage12_head_sha": "0" * 40,
        },
    }
    write_readme(wb, manifest, {"status": "ok-from-loader"}, [])
    readme = {row[0].value: row[1].value for row in wb["README"].iter_rows(min_row=2, max_col=2)}
    assert "DIAGNOSTIC ONLY" in readme["Workbook build provenance"]
    assert "Stage 9A manifest is retained only" in readme["Workbook build provenance"]
    assert readme["Pinned source Run IDs"] == (
        "Stage 7: 37225915321; Stage 10: 37709936740; Stage 12: 37964656154"
    )
    assert readme["Pinned source HEAD SHAs"].startswith("Stage 7: ")
    assert readme["Dataset status"] == "DIAGNOSTIC_ONLY_NOT_FINAL"
    assert readme["Final acceptance status"] == "NOT_FINAL_BLOCKED"
    assert "Official Run ID" not in readme
    wb.close()


def test_adaptive_mode_rejects_rng_audit_map_tampering():
    data = adaptive_data()
    row = data["trials_by_scenario"]["C070_G100"][0]
    row["rng_checksum"] = 987654321987
    with pytest.raises(RuntimeError, match="ADAPTIVE_RNG_AUDIT_MAP_SHA_MISMATCH"):
        _validate_mode_and_identity(data, "adaptive")


def test_adaptive_mode_rejects_duplicate_rng_audit_checksums():
    data = adaptive_data()
    rows = data["trials_by_scenario"]["C070_G100"]
    rows[1]["rng_checksum"] = rows[0]["rng_checksum"]
    # Recompute the audit-map hash to prove the independent duplicate detector
    # rejects reuse even when the manifest itself is internally consistent.
    ordered = [
        f"{scenario}|{row['trial_id']}|{row['rng_checksum']}|{row['rng_draw_count']}\n"
        for scenario in EXPECTED_SCENARIOS
        for row in data["trials_by_scenario"][scenario]
    ]
    data["manifest"]["rng_audit_map_sha256"] = hashlib.sha256(
        "".join(ordered).encode("utf-8")
    ).hexdigest()
    with pytest.raises(RuntimeError, match="DUPLICATE_SEED_DETECTED"):
        _validate_mode_and_identity(data, "adaptive")


def test_adaptive_mode_rejects_stochastic_trial_with_empty_rng_audit():
    data = adaptive_data()
    row = data["trials_by_scenario"]["C070_G100"][0]
    row["rng_checksum"] = 0
    row["rng_draw_count"] = 0
    ordered = [
        f"{scenario}|{row['trial_id']}|{row['rng_checksum']}|{row['rng_draw_count']}\\n"
        for scenario in EXPECTED_SCENARIOS
        for row in data["trials_by_scenario"][scenario]
    ]
    data["manifest"]["rng_audit_map_sha256"] = hashlib.sha256(
        "".join(ordered).encode("utf-8")
    ).hexdigest()
    with pytest.raises(RuntimeError, match="ADAPTIVE_STOCHASTIC_RNG_AUDIT_EMPTY"):
        _validate_mode_and_identity(data, "adaptive")
