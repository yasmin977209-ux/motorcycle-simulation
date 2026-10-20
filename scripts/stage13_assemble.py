"""Stage 13 workbook assembly boundary and fail-closed acceptance helpers.

This module contains no production-data acquisition or simulation code. The
caller supplies already-loaded evidence and the established worksheet helpers.
"""
from __future__ import annotations

import hashlib
from pathlib import Path
from typing import Any, Mapping

from openpyxl import load_workbook


REQUIRED_REPRESENTATIVE_ROLES = ("P10", "P50", "P90", "LossCase", "MaxCase")


def validate_role_payload_mapping(
    role_to_payload_labels: Mapping[str, str | None] | None,
) -> dict[str, str | None]:
    """Require explicit payload labels for present roles; MaxCase may be pending."""
    if role_to_payload_labels is None:
        raise ValueError("STAGE13_ROLE_TO_PAYLOAD_MAPPING_UNRESOLVED")
    if not isinstance(role_to_payload_labels, Mapping):
        raise ValueError("STAGE13_ROLE_TO_PAYLOAD_MAPPING_INVALID_TYPE")

    required = set(REQUIRED_REPRESENTATIVE_ROLES)
    required_without_maxcase = required - {"MaxCase"}
    actual = set(role_to_payload_labels)
    if actual not in (required, required_without_maxcase):
        missing = sorted(required - actual)
        extra = sorted(actual - required)
        raise ValueError(
            f"STAGE13_ROLE_TO_PAYLOAD_MAPPING_KEYS_INVALID: missing={missing}, extra={extra}"
        )

    result: dict[str, str | None] = {}
    for role in REQUIRED_REPRESENTATIVE_ROLES:
        value = role_to_payload_labels.get(role)
        if role == "MaxCase" and (role not in role_to_payload_labels or value is None):
            result[role] = None
            continue
        if type(value) is not str or not value.strip():
            raise ValueError(f"STAGE13_ROLE_TO_PAYLOAD_MAPPING_VALUE_INVALID: {role}")
        result[role] = value

    values = [value for value in result.values() if value is not None]
    if len(set(values)) != len(values):
        raise ValueError("STAGE13_ROLE_TO_PAYLOAD_MAPPING_LABEL_DUPLICATE")
    return result


def validate_representative_assignments(
    assignments: Mapping[str, Any],
    *,
    expected_scenarios: list[str],
    trial_rows_by_scenario: Mapping[str, list[Mapping[str, Any]]] | None = None,
) -> None:
    """Validate scenario/role identity; only an unassigned MaxCase may be pending."""
    expected_scenario_set = set(expected_scenarios)
    actual_scenario_set = set(assignments)
    if actual_scenario_set != expected_scenario_set:
        raise ValueError(
            "STAGE13_REPRESENTATIVE_SCENARIO_KEYS_INVALID: "
            f"missing={sorted(expected_scenario_set - actual_scenario_set)}, "
            f"extra={sorted(actual_scenario_set - expected_scenario_set)}"
        )

    required_roles = set(REQUIRED_REPRESENTATIVE_ROLES)
    required_without_maxcase = required_roles - {"MaxCase"}
    for scenario in expected_scenarios:
        role_map = assignments[scenario]
        if not isinstance(role_map, Mapping):
            raise ValueError(f"STAGE13_REPRESENTATIVE_ROLE_MAP_INVALID: {scenario}")
        actual_roles = set(role_map)
        if actual_roles not in (required_roles, required_without_maxcase):
            raise ValueError(
                f"STAGE13_REPRESENTATIVE_ROLE_KEYS_INVALID: {scenario}; "
                f"missing={sorted(required_roles - actual_roles)}, "
                f"extra={sorted(actual_roles - required_roles)}"
            )
        if "MaxCase" in role_map and role_map["MaxCase"] is None:
            pass
        elif "MaxCase" not in role_map:
            pass
        for role in actual_roles:
            trial_id = role_map[role]
            if role == "MaxCase" and trial_id is None:
                continue
            if type(trial_id) is not int or trial_id < 1:
                raise ValueError(
                    f"STAGE13_REPRESENTATIVE_TRIAL_ID_INVALID: {scenario}/{role}"
                )

    if trial_rows_by_scenario is None:
        return
    if set(trial_rows_by_scenario) != expected_scenario_set:
        raise ValueError("STAGE13_CANONICAL_TRIAL_SCENARIO_KEYS_INVALID")
    for scenario in expected_scenarios:
        counts: dict[int, int] = {}
        for row in trial_rows_by_scenario[scenario]:
            trial_id = row.get("trial_id")
            if type(trial_id) is int:
                counts[trial_id] = counts.get(trial_id, 0) + 1
        for role, trial_id in assignments[scenario].items():
            if role == "MaxCase" and trial_id is None:
                continue
            if counts.get(trial_id, 0) != 1:
                raise ValueError(
                    f"STAGE13_REPRESENTATIVE_TRIAL_BINDING_INVALID: "
                    f"{scenario}/{role}/trial_id={trial_id}/count={counts.get(trial_id, 0)}"
                )


def validate_representative_payload(
    scenario: str,
    role: str,
    assigned_trial_id: int,
    payload: Mapping[str, Any],
) -> None:
    """Require exact scenario/trial identity in a representative's summary and daily rows."""
    if type(assigned_trial_id) is not int or assigned_trial_id < 1:
        raise ValueError(f"STAGE13_REPRESENTATIVE_TRIAL_ID_INVALID: {scenario}/{role}")
    summary = payload.get("summary")
    if not isinstance(summary, Mapping):
        raise ValueError(f"STAGE13_REPRESENTATIVE_SUMMARY_MISSING: {scenario}/{role}")
    if type(summary.get("scenario_id")) is not str or summary.get("scenario_id") != scenario:
        raise ValueError(f"STAGE13_REPRESENTATIVE_SCENARIO_MISMATCH: {scenario}/{role}")
    if type(summary.get("trial_id")) is not int or summary.get("trial_id") != assigned_trial_id:
        raise ValueError(f"STAGE13_REPRESENTATIVE_TRIAL_MISMATCH: {scenario}/{role}")

    daily = payload.get("daily")
    if not isinstance(daily, list):
        raise ValueError(f"STAGE13_REPRESENTATIVE_DAILY_MISSING: {scenario}/{role}")
    if not daily:
        raise ValueError(f"STAGE13_REPRESENTATIVE_DAILY_EMPTY: {scenario}/{role}")
    for index, row in enumerate(daily):
        if not isinstance(row, Mapping):
            raise ValueError(f"STAGE13_REPRESENTATIVE_DAILY_ROW_INVALID: {scenario}/{role}/{index}")
        if type(row.get("trial_id")) is not int or row.get("trial_id") != assigned_trial_id:
            raise ValueError(
                f"STAGE13_REPRESENTATIVE_DAILY_TRIAL_MISMATCH: {scenario}/{role}/{index}"
            )


def expected_sheet_names(
    global_sheets: list[str],
    scenarios: list[str],
    suffixes: list[str],
) -> list[str]:
    """Produce the exact, ordered 185-sheet contract without changing its topology."""
    return list(global_sheets) + [
        f"{scenario}{suffix}" for scenario in scenarios for suffix in suffixes
    ]


def audit_saved_workbook(
    path: str | Path,
    expected_names: list[str],
) -> dict[str, Any]:
    """Read the saved XLSX back and verify ordered sheets, formulas, and error cells."""
    path = Path(path)
    if not path.is_file():
        raise FileNotFoundError(f"STAGE13_OUTPUT_FILE_MISSING: {path}")
    wb = load_workbook(path, data_only=False, read_only=False)
    try:
        names = list(wb.sheetnames)
        error_cells: list[str] = []
        formula_cells: list[str] = []
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    if cell.data_type == "e":
                        error_cells.append(f"{ws.title}!{cell.coordinate}")
                    if cell.data_type == "f":
                        formula_cells.append(f"{ws.title}!{cell.coordinate}")
        audit = {
            "sheet_count": len(names),
            "sheet_names_exact": names == expected_names,
            "expected_sheet_count": len(expected_names),
            "error_cells": error_cells,
            "formula_cells": formula_cells,
        }
    finally:
        wb.close()
    if not audit["sheet_names_exact"]:
        raise RuntimeError(
            f"STAGE13_SHEET_INVENTORY_MISMATCH: expected={len(expected_names)} "
            f"actual={audit['sheet_count']}"
        )
    if audit["error_cells"]:
        raise RuntimeError(f"STAGE13_EXCEL_ERROR_CELLS: {audit['error_cells'][:20]}")
    if audit["formula_cells"]:
        raise RuntimeError(f"STAGE13_FORMULA_CELLS: {audit['formula_cells'][:20]}")
    audit["file_bytes"] = path.stat().st_size
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    audit["sha256"] = digest.hexdigest()
    return audit


def assemble_workbook(
    *,
    builder: Any,
    data: Mapping[str, Any],
    output_path: str | Path,
    role_to_payload_labels: Mapping[str, str],
    mode: str = "canonical",
) -> dict[str, Any]:
    """Build sheets in canonical order using injected helpers, then audit the saved XLSX."""
    from openpyxl import Workbook

    role_map = validate_role_payload_mapping(role_to_payload_labels)
    manifest = data["manifest"]
    all_rows = data["trials_by_scenario"]
    assignments = data["representative_trial_ids"]
    validation = data.get("validation", {"status": "not-provided", "tests": []})
    errors = data.get("errors", [])
    scenarios = list(builder.EXPECTED_SCENARIOS)

    validate_representative_assignments(
        assignments,
        expected_scenarios=scenarios,
        trial_rows_by_scenario=all_rows,
    )
    wb = Workbook()
    builder.write_readme(wb, manifest, validation, [])
    builder.write_constants(wb, manifest)
    summaries = builder.build_scenario_summary(all_rows)
    builder.write_scenario_matrix(wb, summaries, manifest)
    builder.write_heatmap(wb, summaries)
    builder.write_sensitivity(wb, summaries)
    builder.write_validation(wb, validation)
    builder.write_errors(wb, errors)
    builder.write_final_closure(wb, all_rows)
    builder.write_daily_distribution(wb, Path(data["dataset_root"]))
    builder.write_representatives(
        wb, manifest, data["representatives_by_scenario"], assignments, role_map
    )

    for scenario in scenarios:
        p50_trial_id = assignments[scenario]["P50"]
        p50_label = role_map["P50"]
        payload = data["representatives_by_scenario"][scenario]["P50"]
        validate_representative_payload(scenario, "P50", p50_trial_id, payload)
        builder.write_stats_sheet(
            wb, scenario, all_rows[scenario], int(manifest["total_capital"])
        )
        builder.write_scenario_detail(wb, scenario, payload)

    names_expected = expected_sheet_names(
        list(builder.GLOBAL_SHEETS), scenarios, list(builder.SUFFIXES)
    )
    readme = wb["README"]
    for row in readme.iter_rows(min_col=1, max_col=2):
        label = row[0].value
        if label == "Excel sheets":
            row[1].value = len(wb.sheetnames)
        elif label in {"Excel error count", "Excel error cells"}:
            row[1].value = 0

    if list(wb.sheetnames) != names_expected:
        actual = list(wb.sheetnames)
        wb.close()
        raise RuntimeError(
            f"STAGE13_PRE_SAVE_SHEET_INVENTORY_MISMATCH: "
            f"expected={len(names_expected)} actual={len(actual)}"
        )

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    precision = builder.save_workbook_preserving_stats_precision(wb, output_path)
    wb.close()
    audit = audit_saved_workbook(output_path, names_expected)
    if (audit["sheet_count"], audit["error_cells"], audit["formula_cells"]) != (
        185, [], []
    ):
        raise RuntimeError(f"STAGE13_SAVED_XLSX_AUDIT_FAILURE: {audit}")
    pending_maxcase_scenarios = [
        scenario for scenario in scenarios
        if assignments[scenario].get("MaxCase") is None
    ]
    warning_messages = []
    if pending_maxcase_scenarios:
        warning_messages.append(
            "STAGE13_MAXCASE_PENDING_DATA_SOURCE: "
            f"{len(pending_maxcase_scenarios)} scenario role(s) emitted without fabricated payload"
        )
    return {
        "success": True,
        "mode": mode,
        "path": str(output_path),
        "sha256": audit["sha256"],
        "file_bytes": audit["file_bytes"],
        "trial_count": sum(len(rows) for rows in all_rows.values()),
        "sheets_count": audit["sheet_count"],
        "excel_errors": 0,
        "formula_cells": 0,
        "stats_precision_serialization": precision,
        "pending_data_source_role_rows": len(pending_maxcase_scenarios),
        "warnings": warning_messages,
    }
