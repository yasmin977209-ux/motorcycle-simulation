#!/usr/bin/env python3
"""Create a separate corrected XLSX candidate from immutable canonical sources."""
from __future__ import annotations
import hashlib, json, math, os, shutil, subprocess, tarfile, tempfile, zipfile
from pathlib import Path
import numpy as np
import pyarrow.parquet as pq
from openpyxl import load_workbook
from scripts.excel_builder import patch_xlsx_numeric_values

SOURCE_SHA = "f3af4f06fded8b7a4e8c7730b5f2d19201a0a84b"
SOURCE_RUN_ID = 37225915321
WORKBOOK_ARTIFACT_ID = 11507380986
EXPECTED_WORKBOOK_SHA256 = "0fb325fc9e196a534ba8ab7484dcfd0541cefd3a2ca37c3b6f930e3215f62f6d"
CANONICAL_MANIFEST_SHA256 = "892da438955160bb59aff3046271402f57bdd61a4abb35f41db7188f61346ff6"
SCENARIOS = [
    "C100_G100", "C100_G070", "C100_G050", "C100_G030", "C100_G000",
    "C085_G100", "C085_G070", "C085_G050", "C085_G030", "C085_G000",
    "C070_G100", "C070_G070", "C070_G050", "C070_G030", "C070_G000",
    "C050_G100", "C050_G070", "C050_G050", "C050_G030", "C050_G000",
    "C030_G100", "C030_G070", "C030_G050", "C030_G030", "C030_G000",
]
METRICS = {
    "Final_Net_Project_Equity": "final_net_project_equity",
    "Partner1_Final_Entitlement": "partner1_final_entitlement",
    "Partner2_Final_Entitlement": "partner2_final_entitlement",
    "Final_Cash": "final_cash",
}
STAT_COLUMNS = {"MIN": 5, "MAX": 6, "P10": 8, "P50": 10, "P90": 12}
STAT_ORDER = ("MIN", "P10", "P50", "P90", "MAX")


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def download_artifact(artifact_id: int, output: Path) -> None:
    repo = os.environ["GITHUB_REPOSITORY"]
    with output.open("wb") as stream:
        result = subprocess.run(
            ["gh", "api", f"repos/{repo}/actions/artifacts/{artifact_id}/zip"],
            stdout=stream, stderr=subprocess.PIPE, check=False,
        )
    if result.returncode != 0 or not zipfile.is_zipfile(output):
        raise RuntimeError(
            f"ARTIFACT_DOWNLOAD_FAILED id={artifact_id} exit={result.returncode} "
            f"stderr={result.stderr.decode('utf-8', errors='replace')[-1600:]}"
        )


def unique_file(root: Path, basename: str) -> Path:
    candidates = [p for p in root.rglob(basename) if p.is_file()]
    if len(candidates) != 1:
        raise RuntimeError(f"EXPECTED_ONE_{basename}: found={len(candidates)} in {root}")
    return candidates[0]


def extract_zip(zip_path: Path, destination: Path) -> None:
    destination.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(zip_path) as archive:
        root = destination.resolve()
        for info in archive.infolist():
            target = (destination / info.filename).resolve()
            if target != root and root not in target.parents:
                raise RuntimeError(f"ARTIFACT_PATH_TRAVERSAL_REJECTED: {info.filename}")
        archive.extractall(destination)


def copy_expected_tar_members(tar_path: Path, destination: Path, scenario: str) -> None:
    destination.mkdir(parents=True, exist_ok=True)
    needed = ("trial_results.parquet", "scenario_metadata.json", "trial_fingerprints.json")
    with tarfile.open(tar_path, "r:gz") as archive:
        files = {}
        for name in needed:
            matches = [m for m in archive.getmembers() if m.name.removeprefix("./") == name]
            if len(matches) != 1 or not matches[0].isfile():
                raise RuntimeError(f"EXPECTED_ONE_TAR_MEMBER {scenario}/{name}: {len(matches)}")
            stream = archive.extractfile(matches[0])
            if stream is None:
                raise RuntimeError(f"TAR_MEMBER_UNREADABLE {scenario}/{name}")
            data = stream.read()
            (destination / name).write_bytes(data)
            files[name] = hashlib.sha256(data).hexdigest()
    (destination / "source_artifact_tar.sha256").write_text(
        json.dumps(files, indent=2, sort_keys=True) + "\n", encoding="utf-8"
    )


def statistic(values: np.ndarray, name: str) -> int | float:
    if name == "MIN":
        return int(values.min())
    if name == "MAX":
        return int(values.max())
    quantile = {"P10": 0.10, "P50": 0.50, "P90": 0.90}[name]
    return float(np.quantile(values, quantile, method="linear"))


def run() -> dict:
    manifest_path = Path("docs/stage9a/dataset_manifest.json")
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    if manifest.get("global_manifest_sha256") != CANONICAL_MANIFEST_SHA256:
        raise RuntimeError(f"CANONICAL_MANIFEST_SHA_MISMATCH {manifest.get('global_manifest_sha256')}")
    if manifest.get("source_sha") != SOURCE_SHA or int(manifest.get("run_id", -1)) != SOURCE_RUN_ID:
        raise RuntimeError("CANONICAL_MANIFEST_SOURCE_OR_RUN_MISMATCH")
    if int(manifest.get("master_seed", -1)) != 20270101 or int(manifest.get("total_trials", -1)) != 9005:
        raise RuntimeError("CANONICAL_MANIFEST_SEED_OR_TRIAL_COUNT_MISMATCH")
    artifacts = manifest.get("artifact_manifest", {}).get("enriched_artifacts", {})
    if set(artifacts) != set(SCENARIOS):
        raise RuntimeError("CANONICAL_ENRICHED_ARTIFACT_COVERAGE_MISMATCH")

    output_root = Path("results_stage9c")
    output_root.mkdir(parents=True, exist_ok=True)
    package_root = Path(".stage9c_work/dataset")
    package_root.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="stage9c_precision_") as temp:
        tmp = Path(temp)
        workbook_zip = tmp / "published-workbook.zip"
        download_artifact(WORKBOOK_ARTIFACT_ID, workbook_zip)
        workbook_extract = tmp / "workbook"
        extract_zip(workbook_zip, workbook_extract)
        original = unique_file(workbook_extract, "Full_Results_EN.xlsx")
        original_sha = file_sha256(original)
        if original_sha != EXPECTED_WORKBOOK_SHA256:
            raise RuntimeError(f"PUBLISHED_WORKBOOK_SHA_MISMATCH actual={original_sha} expected={EXPECTED_WORKBOOK_SHA256}")
        candidate = output_root / "Full_Results_EN.xlsx"
        shutil.copyfile(original, candidate)

        source_evidence = {}
        target_values: dict[str, dict[str, int | float]] = {}
        # Read-only workbook access: use the published metric labels to locate cells.
        workbook = load_workbook(original, data_only=True, read_only=True)
        try:
            for scenario in SCENARIOS:
                artifact_id = int(artifacts[scenario])
                zip_path = tmp / f"{scenario}.zip"
                download_artifact(artifact_id, zip_path)
                extract_root = tmp / f"extract_{scenario}"
                extract_zip(zip_path, extract_root)
                tar_path = unique_file(extract_root, "*.tar.gz") if False else None
                tarballs = [p for p in extract_root.rglob("*.tar.gz") if p.is_file()]
                if len(tarballs) != 1:
                    raise RuntimeError(f"ENRICHED_TARBALL_COUNT {scenario}: {len(tarballs)}")
                dataset_dir = package_root / scenario
                copy_expected_tar_members(tarballs[0], dataset_dir, scenario)

                metadata = json.loads((dataset_dir / "scenario_metadata.json").read_text(encoding="utf-8"))
                fingerprints = json.loads((dataset_dir / "trial_fingerprints.json").read_text(encoding="utf-8"))
                count = 1 if scenario.startswith("C100_") else 450
                expected_ids = [1] if count == 1 else list(range(1, 451))
                if metadata.get("source_sha") != SOURCE_SHA or str(metadata.get("run_id")) != str(SOURCE_RUN_ID):
                    raise RuntimeError(f"SCENARIO_SOURCE_OR_RUN_MISMATCH {scenario}")
                if metadata.get("scenario_id", metadata.get("scenario")) != scenario:
                    raise RuntimeError(f"SCENARIO_ID_MISMATCH {scenario}")
                if int(metadata.get("trial_count", -1)) != count:
                    raise RuntimeError(f"SCENARIO_TRIAL_COUNT_MISMATCH {scenario}")
                expected_collection = int(scenario[1:4]) / 100
                expected_recovery = int(scenario.split("_G", 1)[1])
                if float(metadata.get("collection_probability", -1)) != expected_collection:
                    raise RuntimeError(f"COLLECTION_PARAMETER_MISMATCH {scenario}")
                if int(metadata.get("recovery_rate_pct", -1)) != expected_recovery:
                    raise RuntimeError(f"RECOVERY_PARAMETER_MISMATCH {scenario}")
                fingerprint_map = fingerprints.get("fingerprints")
                if not isinstance(fingerprint_map, dict) or sorted(int(k) for k in fingerprint_map) != expected_ids:
                    raise RuntimeError(f"TRIAL_FINGERPRINT_COVERAGE_MISMATCH {scenario}")
                if len(fingerprint_map) != count or any(len(str(v)) != 64 for v in fingerprint_map.values()):
                    raise RuntimeError(f"TRIAL_FINGERPRINT_SHAPE_MISMATCH {scenario}")

                parquet_path = dataset_dir / "trial_results.parquet"
                table = pq.read_table(parquet_path)
                rows = table.to_pylist()
                if len(rows) != count or not {"scenario_id", "trial_id", *METRICS.values()}.issubset(table.column_names):
                    raise RuntimeError(f"TRIAL_RESULTS_SHAPE_MISMATCH {scenario}")
                ids = [int(row["trial_id"]) for row in rows]
                if sorted(ids) != expected_ids or len(set(ids)) != count:
                    raise RuntimeError(f"TRIAL_RESULT_ID_MISMATCH {scenario}")
                if any(row["scenario_id"] != scenario for row in rows):
                    raise RuntimeError(f"TRIAL_RESULT_SCENARIO_MISMATCH {scenario}")
                for row in rows:
                    if row.get("source_sha") not in (None, SOURCE_SHA):
                        raise RuntimeError(f"TRIAL_RESULT_SOURCE_SHA_MISMATCH {scenario}/{row['trial_id']}")
                    for field in METRICS.values():
                        value = row[field]
                        if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(float(value)) or int(value) != value:
                            raise RuntimeError(f"INVALID_METRIC {scenario}/{row['trial_id']}/{field}={value!r}")

                worksheet_name = f"{scenario}_STATS"
                if worksheet_name not in workbook.sheetnames:
                    raise RuntimeError(f"STATS_SHEET_MISSING {worksheet_name}")
                ws = workbook[worksheet_name]
                metric_rows = {}
                for row_idx in range(2, ws.max_row + 1):
                    label = ws.cell(row_idx, 1).value
                    if label is not None:
                        metric_rows[str(label)] = row_idx
                for label, field in METRICS.items():
                    if label not in metric_rows:
                        raise RuntimeError(f"STATS_METRIC_ROW_MISSING {scenario}/{label}")
                    values = np.asarray([int(row[field]) for row in rows], dtype=np.float64)
                    row_idx = metric_rows[label]
                    for stat_name in STAT_ORDER:
                        col_idx = STAT_COLUMNS[stat_name]
                        target_values.setdefault(worksheet_name, {})[
                            ws.cell(row_idx, col_idx).coordinate
                        ] = statistic(values, stat_name)
                source_evidence[scenario] = {
                    "artifact_id": artifact_id,
                    "artifact_zip_sha256": file_sha256(zip_path),
                    "tarball_sha256": file_sha256(tarballs[0]),
                    "trial_results_sha256": file_sha256(parquet_path),
                    "scenario_metadata_sha256": file_sha256(dataset_dir / "scenario_metadata.json"),
                    "trial_fingerprints_sha256": file_sha256(dataset_dir / "trial_fingerprints.json"),
                    "source_sha": SOURCE_SHA,
                    "source_run_id": SOURCE_RUN_ID,
                    "master_seed": 20270101,
                    "scenario_trial_count": count,
                }
        finally:
            workbook.close()
        target_count = sum(len(cells) for cells in target_values.values())
        if target_count != 500:
            raise RuntimeError(f"STAT_TARGET_COVERAGE_MISMATCH targets={target_count} expected=500")
        patch_result = patch_xlsx_numeric_values(candidate, target_values)

        # Independent post-write verification over all expected cell values.
        check_wb = load_workbook(candidate, data_only=True, read_only=True)
        verified = 0
        postwrite_errors = []
        try:
            for sheet_name, cells in target_values.items():
                ws = check_wb[sheet_name]
                for cell_ref, expected_value in cells.items():
                    actual = ws[cell_ref].value
                    if isinstance(actual, bool) or not isinstance(actual, (int, float)) or actual != expected_value:
                        postwrite_errors.append({
                            "sheet": sheet_name, "cell": cell_ref,
                            "expected": expected_value, "actual": actual,
                        })
                    verified += 1
        finally:
            check_wb.close()
        report = {
            "status": "PASS" if verified == 500 and not postwrite_errors and patch_result["target_cells_verified"] == 500 else "FAIL",
            "canonical_manifest_global_sha256": CANONICAL_MANIFEST_SHA256,
            "canonical_source_sha": SOURCE_SHA,
            "stage7_source_run_id": SOURCE_RUN_ID,
            "canonical_trial_result_count": 9005,
            "published_workbook_artifact_id": WORKBOOK_ARTIFACT_ID,
            "published_workbook_sha256": original_sha,
            "published_workbook_sha256_expected": EXPECTED_WORKBOOK_SHA256,
            "candidate_workbook_path": str(candidate),
            "candidate_workbook_sha256": file_sha256(candidate),
            "statistics_target_cells": target_count,
            "target_cells_verified_after_write": verified,
            "precision_patch_result": patch_result,
            "postwrite_errors": postwrite_errors,
            "metric_statistics": "MIN, P10, P50, P90, MAX; NumPy quantile method=linear; exact equality; no tolerance/rounding",
            "source_evidence": source_evidence,
            "historical_dataset_modified": False,
            "original_published_workbook_modified": False,
            "trial_execution": "NONE",
        }
        report_path = output_root / "precision_rebuild_report.json"
        report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2, sort_keys=True) + "\n", encoding="utf-8")
        print(json.dumps({
            "status": report["status"], "target_cells": target_count,
            "target_cells_verified_after_write": verified,
            "cells_patched": patch_result["cells_patched"],
            "published_workbook_sha256": original_sha,
            "candidate_workbook_sha256": report["candidate_workbook_sha256"],
            "postwrite_error_count": len(postwrite_errors),
            "report": str(report_path),
        }, ensure_ascii=False, indent=2))
        return report


if __name__ == "__main__":
    result = run()
    if result.get("status") != "PASS":
        raise SystemExit(1)
