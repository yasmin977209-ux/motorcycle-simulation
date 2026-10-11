#!/usr/bin/env python3
"""Read-only full Stage 9C 20-point Parquet-vs-XLSX audit; no simulation or trial execution."""
from __future__ import annotations

import hashlib
import json
import math
import os
import shutil
import subprocess
import tarfile
import tempfile
import zipfile
from pathlib import Path
from typing import Any

import numpy as np
import pyarrow.parquet as pq
from openpyxl import load_workbook

SOURCE_SHA = "f3af4f06fded8b7a4e8c7730b5f2d19201a0a84b"
SOURCE_RUN_ID = 37225915321
CANONICAL_MANIFEST_SHA256 = "892da438955160bb59aff3046271402f57bdd61a4abb35f41db7188f61346ff6"
EXPECTED_WORKBOOK_SHA256 = "0fb325fc9e196a534ba8ab7484dcfd0541cefd3a2ca37c3b6f930e3215f62f6d"
WORKBOOK_ARTIFACT_ID = 11507380986
SCENARIOS = [
    "C100_G100", "C100_G070", "C100_G050", "C100_G030", "C100_G000",
    "C085_G100", "C085_G070", "C085_G050", "C085_G030", "C085_G000",
    "C070_G100", "C070_G070", "C070_G050", "C070_G030", "C070_G000",
    "C050_G100", "C050_G070", "C050_G050", "C050_G030", "C050_G000",
    "C030_G100", "C030_G070", "C030_G050", "C030_G030", "C030_G000",
]
# Existing published enrichment artifacts from Stage 9A runs. These contain no new trial executions.
ARTIFACTS = {
    "C030_G000": 11498686202,
    "C030_G030": 11497409735,
    "C030_G050": 11498069095,
    "C030_G070": 11498204610,
    "C030_G100": 11498980727,
    "C050_G000": 11503510055,
    "C050_G030": 11499677773,
    "C050_G050": 11502991608,
    "C050_G070": 11503357996,
    "C050_G100": 11503961628,
    "C070_G000": 11500301100,
    "C070_G030": 11499331345,
    "C070_G050": 11498901762,
    "C070_G070": 11499668344,
    "C070_G100": 11499938620,
    "C085_G000": 11499442661,
    "C085_G030": 11499958237,
    "C085_G050": 11498784380,
    "C085_G070": 11498798098,
    "C085_G100": 11499553926,
    "C100_G000": 11496939782,
    "C100_G030": 11497898524,
    "C100_G050": 11497164787,
    "C100_G070": 11497369182,
    "C100_G100": 11498636066,
}
METRICS = {
    "Final_Net_Project_Equity": "final_net_project_equity",
    "Partner1_Final_Entitlement": "partner1_final_entitlement",
    "Partner2_Final_Entitlement": "partner2_final_entitlement",
    "Final_Cash": "final_cash",
}
# 1-based Excel columns, copied from the published Stage 13 assembly/audit implementation.
STAT_COLUMNS = {"MIN": 5, "P10": 8, "P50": 10, "P90": 12, "MAX": 6}
STAT_ORDER = ("MIN", "P10", "P50", "P90", "MAX")
STAT_LABELS = ("MIN", "P10", "P50", "P90", "MAX")


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def gh_download_artifact(artifact_id: int, dest: Path) -> None:
    repo = os.environ["GITHUB_REPOSITORY"]
    dest.parent.mkdir(parents=True, exist_ok=True)
    with dest.open("wb") as stream:
        result = subprocess.run(
            ["gh", "api", f"repos/{repo}/actions/artifacts/{artifact_id}/zip"],
            stdout=stream, stderr=subprocess.PIPE, text=False, check=False,
        )
    if result.returncode != 0:
        raise RuntimeError(
            f"ARTIFACT_DOWNLOAD_FAILED id={artifact_id} exit={result.returncode} "
            f"stderr={result.stderr.decode('utf-8', errors='replace')[-2000:]}"
        )
    if dest.stat().st_size == 0 or not zipfile.is_zipfile(dest):
        raise RuntimeError(f"ARTIFACT_ZIP_INVALID id={artifact_id} bytes={dest.stat().st_size}")


def extract_artifact_zip(zip_path: Path, dest: Path) -> None:
    dest.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(zip_path) as archive:
        archive.extractall(dest)
    tar_paths = list(dest.rglob("*.tar.gz")) + list(dest.rglob("*.tgz"))
    for tar_path in tar_paths:
        out = tar_path.parent / (tar_path.name.rsplit(".tar.gz", 1)[0] if tar_path.name.endswith(".tar.gz") else tar_path.stem)
        out.mkdir(parents=True, exist_ok=True)
        root = out.resolve()
        with tarfile.open(tar_path, "r:gz") as archive:
            for member in archive.getmembers():
                target = (out / member.name).resolve()
                if target != root and root not in target.parents:
                    raise RuntimeError(f"TAR_PATH_TRAVERSAL_REJECTED: {tar_path.name}/{member.name}")
                if member.issym() or member.islnk() or member.isdev():
                    raise RuntimeError(f"UNEXPECTED_TAR_ENTRY: {tar_path.name}/{member.name}")
            archive.extractall(out)


def unique_file(root: Path, basename: str) -> Path:
    found = [path for path in root.rglob(basename) if path.is_file()]
    if len(found) != 1:
        raise RuntimeError(f"EXPECTED_ONE_{basename.upper().replace('.', '_')}: found={len(found)} under={root}")
    return found[0]


def load_json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8"))


def validate_scenario_package(scenario: str, root: Path) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    parquet_path = unique_file(root, "trial_results.parquet")
    metadata_path = unique_file(root, "scenario_metadata.json")
    fingerprints_path = unique_file(root, "trial_fingerprints.json")
    metadata = load_json(metadata_path)
    fingerprints = load_json(fingerprints_path)
    # The actual Stage 9A artifact schema has scenario_metadata.json and
    # trial_fingerprints.json. dataset_identity.json is not part of this package.
    identity_paths = [p for p in root.rglob("dataset_identity.json") if p.is_file()]
    if len(identity_paths) > 1:
        raise RuntimeError(f"MULTIPLE_DATASET_IDENTITY_FILES {scenario}")
    identity = load_json(identity_paths[0]) if identity_paths else {}

    expected_count = 1 if scenario.startswith("C100_") else 450
    expected_ids = [1] if expected_count == 1 else list(range(1, 451))
    metadata_scenario = metadata.get("scenario_id", metadata.get("scenario"))
    fingerprints_scenario = fingerprints.get("scenario_id", fingerprints.get("scenario", metadata_scenario))
    source_values = [metadata.get("source_sha"), fingerprints.get("source_sha")]
    if identity:
        source_values.append(identity.get("source_sha"))
    if metadata_scenario != scenario or fingerprints_scenario != scenario:
        raise RuntimeError(
            f"SCENARIO_IDENTITY_MISMATCH {scenario}: metadata={metadata_scenario}, fingerprints={fingerprints_scenario}"
        )
    if any(value != SOURCE_SHA for value in source_values):
        raise RuntimeError(f"SOURCE_SHA_MISMATCH {scenario}: metadata+fingerprints={source_values}")
    for source_record, label in ((metadata, "scenario_metadata"), (fingerprints, "trial_fingerprints")):
        if str(source_record.get("run_id")) != str(SOURCE_RUN_ID):
            raise RuntimeError(f"{label.upper()}_RUN_ID_MISMATCH {scenario}: {source_record.get('run_id')}")
        if int(source_record.get("trial_count", -1)) != expected_count:
            raise RuntimeError(f"{label.upper()}_TRIAL_COUNT_MISMATCH {scenario}: {source_record.get('trial_count')}")
    reconciliation = metadata.get("fingerprint_reconciliation", {})
    if int(reconciliation.get("expected", -1)) != expected_count or int(reconciliation.get("matched", -1)) != expected_count:
        raise RuntimeError(f"FINGERPRINT_RECONCILIATION_MISMATCH {scenario}: {reconciliation}")
    fingerprint_values = fingerprints.get("fingerprints")
    if not isinstance(fingerprint_values, dict):
        raise RuntimeError(f"TRIAL_FINGERPRINT_MAP_MISSING {scenario}")
    fingerprint_ids = sorted(int(k) for k in fingerprint_values)
    if fingerprint_ids != expected_ids or any(len(str(v)) != 64 for v in fingerprint_values.values()):
        raise RuntimeError(f"TRIAL_FINGERPRINT_COVERAGE_MISMATCH {scenario}: count={len(fingerprint_values)}")

    expected_collection = int(scenario[1:4]) / 100
    expected_recovery = int(scenario.split("_G", 1)[1])
    if float(metadata.get("collection_probability", -1)) != expected_collection:
        raise RuntimeError(f"COLLECTION_PARAMETER_MISMATCH {scenario}: {metadata.get('collection_probability')} expected={expected_collection}")
    if int(metadata.get("recovery_rate_pct", -1)) != expected_recovery:
        raise RuntimeError(f"RECOVERY_PARAMETER_MISMATCH {scenario}: {metadata.get('recovery_rate_pct')} expected={expected_recovery}")
    # The published package omits MASTER_SEED. The original producing workflow
    # at e0b2fab pins MASTER_SEED=20270101 and the engine constants.py does too.
    seed = 20270101

    table = pq.read_table(parquet_path)
    rows = table.to_pylist()
    if len(rows) != expected_count:
        raise RuntimeError(f"PARQUET_ROW_COUNT_MISMATCH {scenario}: {len(rows)} expected={expected_count}")
    if not {"scenario_id", "trial_id", *METRICS.values()}.issubset(set(table.column_names)):
        missing = sorted({"scenario_id", "trial_id", *METRICS.values()} - set(table.column_names))
        raise RuntimeError(f"PARQUET_REQUIRED_FIELDS_MISSING {scenario}: {missing}")
    ids = [int(row["trial_id"]) for row in rows]
    if sorted(ids) != expected_ids or len(set(ids)) != expected_count:
        raise RuntimeError(f"PARQUET_TRIAL_IDS_MISMATCH {scenario}: count={len(ids)} min={min(ids, default=None)} max={max(ids, default=None)}")
    if any(row["scenario_id"] != scenario for row in rows):
        raise RuntimeError(f"PARQUET_SCENARIO_COLUMN_MISMATCH {scenario}")
    for row in rows:
        for field in METRICS.values():
            value = row[field]
            if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(float(value)):
                raise RuntimeError(f"NON_NUMERIC_METRIC {scenario}/{row.get('trial_id')}/{field}: {value!r}")
            if int(value) != value:
                raise RuntimeError(f"NON_INTEGER_MONEY_METRIC {scenario}/{row.get('trial_id')}/{field}: {value!r}")

    package_evidence = {
        "scenario": scenario,
        "artifact_id": ARTIFACTS[scenario],
        "source_run_id": SOURCE_RUN_ID,
        "source_sha": SOURCE_SHA,
        "master_seed": seed,
        "master_seed_source": "Stage7 workflow commit e0b2fab6710833774937d778097363d7a8ee8e59 (workflow blob f6235989e152df48dce86dd059f66f8104c227e6) sets MASTER_SEED=20270101; constants.py blob 88b3e7b5a509f8486597ac08dfc5acaf80d4d39b defines the same value",
        "trial_count": len(rows),
        "trial_ids_first": min(ids),
        "trial_ids_last": max(ids),
        "trial_result_sha256": sha256_file(parquet_path),
        "scenario_metadata_sha256": sha256_file(metadata_path),
        "dataset_identity_file_present": bool(identity_paths),
        "dataset_identity_sha256": sha256_file(identity_paths[0]) if identity_paths else None,
        "trial_fingerprints_sha256": sha256_file(fingerprints_path),
        "trial_fingerprint_record_count": len(fingerprint_values),
        "official_block_fingerprints": metadata.get("official_block_fingerprints"),
        "schema_version": metadata.get("schema_version"),
        "schema": [{"name": field.name, "type": str(field.type), "nullable": field.nullable} for field in table.schema],
    }
    return rows, package_evidence

def compute_stat(values: np.ndarray, name: str) -> int | float:
    if name == "MIN":
        return int(values.min())
    if name == "MAX":
        return int(values.max())
    q = {"P10": 0.10, "P50": 0.50, "P90": 0.90}[name]
    return float(np.quantile(values, q, method="linear"))


def audit_workbook_structure(workbook_path: Path) -> dict[str, Any]:
    wb = load_workbook(workbook_path, data_only=False, read_only=True)
    errors: list[str] = []
    formulas: list[str] = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.data_type == "e":
                    errors.append(f"{ws.title}!{cell.coordinate}:{cell.value}")
                if cell.data_type == "f":
                    formulas.append(f"{ws.title}!{cell.coordinate}")
    audit = {
        "sheet_count": len(wb.sheetnames),
        "sheet_names": list(wb.sheetnames),
        "excel_error_cells": errors,
        "formula_cells": formulas,
        "stats_sheets_present": all(f"{scenario}_STATS" in wb.sheetnames for scenario in SCENARIOS),
    }
    wb.close()
    return audit


def run() -> dict[str, Any]:
    report: dict[str, Any] = {
        "gate_id": "STAGE9C_CANONICAL_20_POINT_PARQUET_VS_STATS",
        "status": "PENDING",
        "scope": {
            "canonical_scenario_count": 25,
            "expected_trials_by_scenario": {"C100": 1, "non_C100": 450},
            "expected_trial_results": 9005,
            "metrics": list(METRICS.keys()),
            "statistics": list(STAT_ORDER),
            "points_per_scenario": 20,
            "expected_points_total": 500,
        },
        "provenance": {
            "historical_stage7_run_id": SOURCE_RUN_ID,
            "source_sha": SOURCE_SHA,
            "master_seed": 20270101,
            "stage9a_build_run_id": 37671448546,
            "stage9a_workbook_artifact_id": WORKBOOK_ARTIFACT_ID,
            "stage9a_workbook_expected_sha256": EXPECTED_WORKBOOK_SHA256,
            "canonical_manifest_sha256_from_published_stage9a_report": CANONICAL_MANIFEST_SHA256,
        "master_seed": 20270101,
        "master_seed_source": "original Stage7 workflow at e0b2fab6710833774937d778097363d7a8ee8e59 sets MASTER_SEED=20270101; constants.py at SOURCE_SHA also defines 20270101",
            "enrichment_artifact_ids": ARTIFACTS,
            "audit_workflow_run_id": os.environ.get("GITHUB_RUN_ID"),
            "audit_workflow_head_sha": os.environ.get("GITHUB_SHA"),
        },
        "gates": {},
        "scenarios": {},
        "comparisons": [],
        "fatal_errors": [],
    }
    output = Path("results_stage9c")
    output.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="stage9c_audit_") as tmp:
        tmp_root = Path(tmp)
        try:
            workbook_zip = tmp_root / "stage9a-final-output.zip"
            gh_download_artifact(WORKBOOK_ARTIFACT_ID, workbook_zip)
            workbook_root = tmp_root / "workbook_artifact"
            extract_artifact_zip(workbook_zip, workbook_root)
            workbook_path = unique_file(workbook_root, "Full_Results_EN.xlsx")
            workbook_sha = sha256_file(workbook_path)
            report["provenance"]["workbook_actual_sha256"] = workbook_sha
            report["provenance"]["workbook_artifact_zip_sha256"] = sha256_file(workbook_zip)
            report["gates"]["canonical_workbook_sha256"] = workbook_sha == EXPECTED_WORKBOOK_SHA256
            workbook_audit = audit_workbook_structure(workbook_path)
            report["workbook_audit"] = workbook_audit
            report["gates"]["workbook_has_185_sheets"] = workbook_audit["sheet_count"] == 185
            report["gates"]["all_25_stats_sheets_present"] = workbook_audit["stats_sheets_present"]
            report["gates"]["zero_excel_error_cells"] = len(workbook_audit["excel_error_cells"]) == 0
            report["gates"]["zero_formula_cells"] = len(workbook_audit["formula_cells"]) == 0

            # Load the published full-scope 9,005-result artifact, not the separate Stage 13 625-result cohort.
            wb = load_workbook(workbook_path, data_only=True, read_only=True)
            result_rows: dict[str, list[dict[str, Any]]] = {}
            all_keys: list[tuple[str, int]] = []
            total_rows = 0
            for index, scenario in enumerate(SCENARIOS, start=1):
                try:
                    zip_path = tmp_root / f"{scenario}.zip"
                    gh_download_artifact(ARTIFACTS[scenario], zip_path)
                    scenario_root = tmp_root / scenario
                    extract_artifact_zip(zip_path, scenario_root)
                    rows, evidence = validate_scenario_package(scenario, scenario_root)
                    result_rows[scenario] = rows
                    all_keys.extend((str(row["scenario_id"]), int(row["trial_id"])) for row in rows)
                    total_rows += len(rows)
                    evidence["artifact_zip_sha256"] = sha256_file(zip_path)
                    report["scenarios"][scenario] = {"status": "PASS", "evidence": evidence}
                    stats_ws = wb[f"{scenario}_STATS"]
                    metric_rows: dict[str, tuple[Any, ...]] = {}
                    for row in stats_ws.iter_rows(min_row=2, values_only=True):
                        if row and row[0] is not None:
                            metric_rows[str(row[0])] = row
                    for metric_label, field in METRICS.items():
                        if metric_label not in metric_rows:
                            raise RuntimeError(f"STATS_METRIC_MISSING {scenario}/{metric_label}")
                        values = np.asarray([int(row[field]) for row in rows], dtype=np.float64)
                        if not np.all(np.isfinite(values)):
                            raise RuntimeError(f"NON_FINITE_PARQUET_VALUES {scenario}/{field}")
                        stats_row = metric_rows[metric_label]
                        for point in STAT_ORDER:
                            actual = stats_row[STAT_COLUMNS[point] - 1]
                            expected = compute_stat(values, point)
                            actual_numeric = (
                                not isinstance(actual, bool)
                                and isinstance(actual, (int, float))
                                and math.isfinite(float(actual))
                            )
                            exact = bool(actual_numeric and actual == expected)
                            report["comparisons"].append({
                                "scenario_id": scenario,
                                "metric": metric_label,
                                "parquet_field": field,
                                "statistic": point,
                                "parquet_expected": expected,
                                "xlsx_actual": actual,
                                "numeric_finite": actual_numeric,
                                "exact_match": exact,
                                "status": "PASS" if exact else "FAIL",
                            })
                except Exception as exc:
                    report["scenarios"][scenario] = {
                        "status": "FAIL",
                        "error": str(exc),
                    }
            wb.close()
            report["gates"]["scenario_packages_25_of_25"] = len(result_rows) == 25 and all(
                report["scenarios"].get(s, {}).get("status") == "PASS" for s in SCENARIOS
            )
            report["gates"]["trial_result_count_9005"] = total_rows == 9005
            report["gates"]["unique_trial_pairs_9005"] = len(all_keys) == 9005 and len(set(all_keys)) == 9005
            report["gates"]["comparison_count_500"] = len(report["comparisons"]) == 500
            report["gates"]["all_500_numeric_exact_matches"] = len(report["comparisons"]) == 500 and all(
                item["numeric_finite"] and item["exact_match"] for item in report["comparisons"]
            )
            report["gates"]["source_identity_all_scenarios"] = (
                len(result_rows) == 25 and all(
                    report["scenarios"].get(s, {}).get("evidence", {}).get("source_sha") == SOURCE_SHA
                    and report["scenarios"].get(s, {}).get("evidence", {}).get("source_run_id") == SOURCE_RUN_ID
                    and report["scenarios"].get(s, {}).get("evidence", {}).get("master_seed") == 20270101
                    for s in SCENARIOS
                )
            )
        except Exception as exc:
            report["fatal_errors"].append(str(exc))
            report["gates"]["audit_pipeline_completed"] = False
        else:
            report["gates"]["audit_pipeline_completed"] = True

    report["status"] = "PASS" if not report["fatal_errors"] and report["gates"] and all(report["gates"].values()) else "FAIL"
    report["canonical_dataset_modified"] = False
    report["simulation_or_trials_executed"] = False
    report["not_gate_a_or_gate_b"] = True
    report_path = output / "stage9c_20point_gate_report.json"
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    lines = [
        "# Stage 9C canonical 20-point Parquet vs _STATS audit",
        "",
        f"Status: {report['status']}",
        "",
        "This is a read-only independent audit of the historical full 9,005-result Stage 9A dataset and its published workbook.",
        "No simulation or trial was run; no historical dataset, metadata, or fingerprint was modified.",
        "",
        "## Gate results",
        "",
        "| Gate | Result |",
        "|---|---|",
    ]
    lines.extend(f"| {name} | {'PASS' if value else 'FAIL'} |" for name, value in report.get("gates", {}).items())
    lines.extend([
        "",
        f"Comparisons recorded: {len(report.get('comparisons', []))}/500.",
        "",
        "## Identity",
        "",
        f"- Stage 7 source run: {SOURCE_RUN_ID}",
        f"- Source SHA: {SOURCE_SHA}",
        f"- Published Stage 9A build run: 37671448546",
        f"- Expected canonical workbook SHA-256: {EXPECTED_WORKBOOK_SHA256}",
        f"- Canonical manifest SHA-256 from published report: {CANONICAL_MANIFEST_SHA256}",
        "- Stage 13 625-trial sample is not used for this full-scope audit.",
        "",
        "## Scenario outcomes",
        "",
        "| Scenario | Result | Trials | Detail |",
        "|---|---|---:|---|",
    ])
    for scenario in SCENARIOS:
        item = report.get("scenarios", {}).get(scenario, {})
        evidence = item.get("evidence", {})
        detail = item.get("error", "")
        lines.append(f"| {scenario} | {item.get('status', 'NOT_RUN')} | {evidence.get('trial_count', 0)} | {detail} |")
    if report.get("fatal_errors"):
        lines.extend(["", "## Fatal errors", ""])
        lines.extend(f"- {error}" for error in report["fatal_errors"])
    (output / "stage9c_20point_gate_report.md").write_text("\n".join(lines) + "\n", encoding="utf-8")
    print(json.dumps({
        "status": report["status"],
        "gates": report["gates"],
        "comparison_count": len(report["comparisons"]),
        "fatal_errors": report["fatal_errors"],
        "report_path": str(report_path),
    }, ensure_ascii=False, indent=2))
    return report


if __name__ == "__main__":
    final_report = run()
    if final_report.get("status") != "PASS":
        raise SystemExit(1)
