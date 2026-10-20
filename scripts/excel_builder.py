from __future__ import annotations

import argparse
import json
import hashlib
import math
import os
import posixpath
import re
import tarfile
import tempfile
import zipfile
import xml.etree.ElementTree as ET
from datetime import datetime, date
from pathlib import Path
from typing import Any

import numpy as np
import pyarrow.parquet as pq
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


EXPECTED_SOURCE_SHA = "f3af4f06fded8b7a4e8c7730b5f2d19201a0a84b"
EXPECTED_SCENARIOS = [
    "C100_G100", "C100_G070", "C100_G050", "C100_G030", "C100_G000",
    "C085_G100", "C085_G070", "C085_G050", "C085_G030", "C085_G000",
    "C070_G100", "C070_G070", "C070_G050", "C070_G030", "C070_G000",
    "C050_G100", "C050_G070", "C050_G050", "C050_G030", "C050_G000",
    "C030_G100", "C030_G070", "C030_G050", "C030_G030", "C030_G000",
]
GLOBAL_SHEETS = [
    "README", "CONSTANTS", "SCENARIO_MATRIX", "HEATMAP", "SENSITIVITY",
    "VALIDATION_SUMMARY", "ERROR_LOG", "FINAL_CLOSURE_SUMMARY",
    "DAILY_DISTRIBUTION", "REPRESENTATIVE_CASES",
]
SUFFIXES = ["_STATS", "_DAILY", "_ASSETS", "_ACCTG", "_PARTNER", "_GUARANTEE", "_CONTRACTS"]
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
RIGHT = Alignment(horizontal="right", vertical="top")
CENTER = Alignment(horizontal="center", vertical="center")
FINANCE_FMT = "#,##0"
PERCENT_FMT = "0.0%"
DATE_FMT = "yyyy-mm-dd"


def _stage13_module():
    try:
        from scripts import stage13_assemble
    except ModuleNotFoundError:
        import stage13_assemble
        return stage13_assemble
    return stage13_assemble


def load_json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8"))


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as fh:
        for chunk in iter(lambda: fh.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def ensure_date(value: Any) -> Any:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        try:
            return date.fromisoformat(value)
        except ValueError:
            return value
    return value


def py_value(value: Any) -> Any:
    if hasattr(value, "as_py"):
        value = value.as_py()
    if isinstance(value, (np.integer,)):
        return int(value)
    if isinstance(value, (np.floating,)):
        return float(value)
    if isinstance(value, np.bool_):
        return bool(value)
    if isinstance(value, (dict, list, tuple)):
        return json.dumps(
            value,
            ensure_ascii=False,
            sort_keys=True,
            separators=(",", ":"),
            default=str,
        )
    return ensure_date(value)


def style_sheet(ws, freeze="A2"):
    ws.sheet_view.rightToLeft = True
    ws.freeze_panes = freeze
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = RIGHT


def write_table(ws, headers, rows, start_row=1, number_formats=None):
    for col, header in enumerate(headers, 1):
        cell = ws.cell(start_row, col, header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER
    for r_index, row in enumerate(rows, start_row + 1):
        for col, value in enumerate(row, 1):
            cell = ws.cell(r_index, col, py_value(value))
            cell.alignment = RIGHT
            if number_formats and headers[col - 1] in number_formats:
                cell.number_format = number_formats[headers[col - 1]]
            if isinstance(cell.value, date):
                cell.number_format = DATE_FMT
    for col in range(1, len(headers) + 1):
        max_len = max(
            len(str(ws.cell(r, col).value or "")) for r in range(start_row, ws.max_row + 1)
        )
        ws.column_dimensions[get_column_letter(col)].width = min(max(max_len + 2, 10), 42)


def load_scenario_trials(dataset_root: Path, scenario: str):
    path = dataset_root / scenario / "trial_results.parquet"
    if not path.exists():
        raise FileNotFoundError(path)
    return pq.read_table(path).to_pylist()


def load_scenario_daily_distribution(dataset_root: Path, scenario: str):
    path = dataset_root / scenario / "daily_distribution.parquet"
    if not path.exists():
        raise FileNotFoundError(path)
    return pq.read_table(path).to_pylist()


def load_rep_payload(dataset_root: Path, scenario: str, label: str) -> dict:
    direct = dataset_root / scenario / "representatives" / f"{label}.json"
    if direct.exists():
        return load_json(direct)
    tar_path = dataset_root / scenario / f"{scenario}.tar.gz"
    if not tar_path.exists():
        raise FileNotFoundError(tar_path)
    target = f"representatives/{label}.json"
    with tarfile.open(tar_path, "r:gz") as tf:
        member = next(
            (m for m in tf.getmembers() if m.name.lstrip("./") == target),
            None,
        )
        if member is None:
            raise FileNotFoundError(f"{tar_path}:{target}")
        handle = tf.extractfile(member)
        if handle is None:
            raise RuntimeError(f"cannot read {tar_path}:{target}")
        return json.load(handle)


def metric_stats(rows, field):
    vals = np.asarray([row[field] for row in rows], dtype=float)
    return [
        float(vals.mean()),
        float(np.median(vals)),
        float(vals.std(ddof=1)) if len(vals) > 1 else 0.0,
        float(vals.min()),
        float(vals.max()),
        float(np.quantile(vals, 0.05, method="linear")),
        float(np.quantile(vals, 0.10, method="linear")),
        float(np.quantile(vals, 0.25, method="linear")),
        float(np.quantile(vals, 0.50, method="linear")),
        float(np.quantile(vals, 0.75, method="linear")),
        float(np.quantile(vals, 0.90, method="linear")),
        float(np.quantile(vals, 0.95, method="linear")),
    ]


def numeric_metrics(rows):
    return [
        ("Final_Net_Project_Equity", "final_net_project_equity"),
        ("Final_Cash", "final_cash"),
        ("Cumulative_Project_Profit", "cumulative_project_profit"),
        ("Operating_Net_Profit", "operating_net_profit"),
        ("Termination_Count", "termination_count"),
        ("Secondary_Cycle_Count", "secondary_cycle_count"),
        ("Owned_Bikes", "owned_bikes"),
        ("Held_Assets", "held_assets"),
        ("Total_Operating_Revenue", "total_operating_revenue"),
        ("Bad_Debt", "bad_debt"),
        ("Guarantee_Recovered", "guarantee_recovered"),
        ("Partner1_Final_Entitlement", "partner1_final_entitlement"),
        ("Partner2_Final_Entitlement", "partner2_final_entitlement"),
    ]


def representative_rows(rep_path: Path, scenario: str):
    rep = load_json(rep_path)["scenarios"][scenario]
    result = {}
    for label, trial_id in rep.items():
        if label not in {"P10", "P50", "P90", "Loss"}:
            continue
        result[label] = int(trial_id)
    return result


def write_readme(wb, manifest, validation, excel_errors):
    ws = wb.active
    ws.title = "README"
    rows = [
        ("Project", "motorcycle-simulation — Motorcycle Rental Simulation — Sana'a"),
        ("Workbook build provenance", f"Stage 9A from official Run {manifest['run_id']}"),
        ("Scenario count", len(EXPECTED_SCENARIOS)),
        ("Total trials", manifest["total_trials"]),
        ("Total representatives", manifest["total_representatives"]),
        ("MASTER_SEED", manifest["master_seed"]),
        ("Source SHA", manifest["source_sha"]),
        ("Official Run ID", manifest["run_id"]),
        ("Official Run HEAD SHA", manifest["run_head_sha"]),
        ("Reference SHA-256", manifest.get("reference_sha256")),
        ("Dataset manifest SHA-256", manifest.get("global_manifest_sha256")),
        ("Dataset status", manifest.get("dataset_status", "verified enrichment")),
        ("Validation status", validation.get("status", "not-provided")),
        ("Excel sheets", len(wb.sheetnames)),
        ("Excel error count", len(excel_errors)),
        ("STATISTICAL_POLICY_LEVEL", "FIXED_LOOK_WITH_CS_PENDING"),
        ("MaxCase", "PENDING_DATA_SOURCE (trial_id=60 for C050_G000; scenario-specific IDs in REPRESENTATIVE_CASES)"),
        ("log_events", "PENDING_APPROVAL; V3 metadata used provisionally"),
        ("Gate A", "BLOCKED"),
        ("LossCase source mapping", "LossCase -> representatives/Loss.json"),
        ("MaxCase pending rows", len(EXPECTED_SCENARIOS)),
        ("Workbook status", "PROVISIONAL — NOT ACCEPTED"),
        ("Release status", "BLOCKED / NOT_AUTHORIZED"),
        ("C050_G000 Artifact override", "V3 Artifact 11503510055 / Run 37662891608 (provisional)"),
        ("Language", "English — technical schema headers and sheet identifiers preserved"),
    ]
    write_table(ws, ["Item", "Value"], rows)
    style_sheet(ws, "A2")


def write_constants(wb, manifest):
    ws = wb.create_sheet("CONSTANTS")
    rows = manifest.get("constants", [])
    if not rows:
        rows = [{"name": k, "value": v, "type": type(v).__name__, "reference_ref": ""} for k, v in manifest.get("constant_map", {}).items()]
    headers = ["name", "value", "type", "reference_ref"]
    write_table(ws, headers, [[r.get(h) for h in headers] for r in rows], number_formats={"value": FINANCE_FMT})
    style_sheet(ws)


def build_scenario_summary(all_rows):
    out = {}
    for scenario, rows in all_rows.items():
        profits = np.asarray([r["cumulative_project_profit"] for r in rows], dtype=float)
        cash = np.asarray([r["final_cash"] for r in rows], dtype=float)
        term = np.asarray([r["termination_count"] for r in rows], dtype=float)
        held = np.asarray([r["held_assets"] for r in rows], dtype=float)
        loss_prob = float(np.mean(profits < 0))
        out[scenario] = {
            "trial_count": len(rows),
            "p10_profit": float(np.quantile(profits, 0.10, method="linear")),
            "p50_profit": float(np.quantile(profits, 0.50, method="linear")),
            "p90_profit": float(np.quantile(profits, 0.90, method="linear")),
            "p10_cash": float(np.quantile(cash, 0.10, method="linear")),
            "p50_cash": float(np.quantile(cash, 0.50, method="linear")),
            "p90_cash": float(np.quantile(cash, 0.90, method="linear")),
            "loss_probability": loss_prob,
            "p50_termination": float(np.quantile(term, 0.50, method="linear")),
            "p90_termination": float(np.quantile(term, 0.90, method="linear")),
            "mean_held_assets": float(held.mean()),
            "mean_profit": float(profits.mean()),
            "mean_final_equity": float(np.mean([r["final_net_project_equity"] for r in rows])),
        }
    return out


def write_scenario_matrix(wb, summaries, manifest):
    ws = wb.create_sheet("SCENARIO_MATRIX")
    headers = [
        "scenario","actual_trials","stability","P10_Profit","P50_Profit","P90_Profit",
        "P10_Final_Cash","P50_Final_Cash","P90_Final_Cash","Probability_of_Loss",
        "P50_Terminations","P90_Terminations","Mean_Held_Assets","notes"
    ]
    rows = []
    for s in EXPECTED_SCENARIOS:
        x = summaries[s]
        rows.append([
            s, x["trial_count"], manifest["scenarios"][s].get("final_status", "stable"), x["p10_profit"], x["p50_profit"], x["p90_profit"],
            x["p10_cash"], x["p50_cash"], x["p90_cash"], x["loss_probability"],
            x["p50_termination"], x["p90_termination"], x["mean_held_assets"], ""
        ])
    write_table(ws, headers, rows, number_formats={
        "P10_Profit":FINANCE_FMT,"P50_Profit":FINANCE_FMT,"P90_Profit":FINANCE_FMT,
        "P10_Final_Cash":FINANCE_FMT,"P50_Final_Cash":FINANCE_FMT,"P90_Final_Cash":FINANCE_FMT,
        "Probability_of_Loss":PERCENT_FMT,
        "Mean_Held_Assets":"0.0",
    })
    style_sheet(ws)


def write_heatmap(wb, summaries):
    ws = wb.create_sheet("HEATMAP")
    collections = [100,85,70,50,30]
    recoveries = [100,70,50,30,0]
    row = 1
    for metric_name, key in [
        ("Mean Cumulative Project Profit","mean_profit"),
        ("P10 Cumulative Project Profit","p10_profit"),
        ("P90 Cumulative Project Profit","p90_profit"),
    ]:
        ws.cell(row,1,metric_name).font = Font(bold=True)
        row += 1
        write_table(
            ws,
            ["Collection \ Guarantee"] + [f"G{g:02d}" for g in recoveries],
            [
                [f"C{c:03d}"] + [summaries[f"C{c:03d}_G{g:03d}"][key] for g in recoveries]
                for c in collections
            ],
            start_row=row,
            number_formats={f"G{g:02d}":FINANCE_FMT for g in recoveries},
        )
        row = ws.max_row + 2
    style_sheet(ws)


def write_sensitivity(wb, summaries):
    ws = wb.create_sheet("SENSITIVITY")
    rows = []
    for c in [100,85,70,50,30]:
        vals = [summaries[f"C{c:03d}_G{g:03d}"]["mean_profit"] for g in [100,70,50,30,0]]
        base = vals[0]
        rows.append(["Recovery effect at fixed collection", c, *vals, *(v-base for v in vals)])
    headers = [
        "dimension","fixed_collection","G100","G70","G50","G30","G00",
        "ΔG100","ΔG70","ΔG50","ΔG30","ΔG00"
    ]
    write_table(ws, headers, rows, number_formats={h:FINANCE_FMT for h in headers[2:]})
    start = ws.max_row + 2
    rows = []
    for g in [100,70,50,30,0]:
        vals = [summaries[f"C{c:03d}_G{g:03d}"]["mean_profit"] for c in [100,85,70,50,30]]
        base = vals[0]
        rows.append(["Collection effect at fixed recovery", g, *vals, *(v-base for v in vals)])
    write_table(ws, headers, rows, start_row=start, number_formats={h:FINANCE_FMT for h in headers[2:]})
    style_sheet(ws)


def write_validation(wb, validation):
    ws = wb.create_sheet("VALIDATION_SUMMARY")
    source = validation.get("source_sha", "")
    run_id = validation.get("run_id", "")
    rows = validation.get("tests", [])
    headers = ["ID","Name","Expected","Actual","Status","Evidence","Source","Run"]
    materialized = []
    for r in rows:
        materialized.append([
            r.get("ID"), r.get("Name"), r.get("Expected"), r.get("Actual"),
            r.get("Status"), r.get("Evidence"), source, run_id
        ])
    write_table(ws, headers, materialized)
    style_sheet(ws)


def write_errors(wb, errors):
    ws = wb.create_sheet("ERROR_LOG")
    rows = errors if errors else [{"severity":"INFO","code":"NO_ERRORS_RECORDED","message":"No errors or warnings were recorded by the enrichment/build pipeline."}]
    headers = ["severity","code","message","source","details"]
    write_table(ws, headers, [[r.get(k) for k in headers] for r in rows])
    style_sheet(ws)


def write_final_closure(wb, all_rows):
    ws = wb.create_sheet("FINAL_CLOSURE_SUMMARY")
    headers = ["scenario","trial_id","final_close_date","final_net_project_equity","final_cash","cumulative_project_profit","termination_count","held_assets"]
    rows = []
    for s in EXPECTED_SCENARIOS:
        for r in all_rows[s]:
            rows.append([s,r["trial_id"],r["final_close_date"],r["final_net_project_equity"],r["final_cash"],r["cumulative_project_profit"],r["termination_count"],r["held_assets"]])
    write_table(ws, headers, rows, number_formats={"final_net_project_equity":FINANCE_FMT,"final_cash":FINANCE_FMT,"cumulative_project_profit":FINANCE_FMT,"termination_count":"#,##0","held_assets":"#,##0"})
    style_sheet(ws)


def write_daily_distribution(wb, dataset_root):
    ws = wb.create_sheet("DAILY_DISTRIBUTION")
    headers = [
        "scenario","date","active_trial_count",
        "P10_Cash","P25_Cash","P50_Cash","P75_Cash","P90_Cash",
        "P10_Net_Equity","P25_Net_Equity","P50_Net_Equity","P75_Net_Equity","P90_Net_Equity",
        "P10_Active_Bikes","P25_Active_Bikes","P50_Active_Bikes","P75_Active_Bikes","P90_Active_Bikes",
        "P10_Owned_Transferred_Bikes","P25_Owned_Transferred_Bikes","P50_Owned_Transferred_Bikes","P75_Owned_Transferred_Bikes","P90_Owned_Transferred_Bikes",
        "P10_Pending_Claims","P25_Pending_Claims","P50_Pending_Claims","P75_Pending_Claims","P90_Pending_Claims"
    ]
    metrics=["Cash","Net_Equity","Active_Bikes","Owned_Transferred_Bikes","Pending_Claims"]
    quantiles=["P10","P25","P50","P75","P90"]
    rows=[]
    for s in EXPECTED_SCENARIOS:
        raw=load_scenario_daily_distribution(dataset_root,s)
        if not raw:
            continue
        if "quantile_name" in raw[0]:
            by_date={}
            for item in raw:
                d=item["date"]
                record=by_date.setdefault(d,{"date":d,"active_trial_count":item.get("active_trial_count")})
                q=item["quantile_name"]
                for metric in metrics:
                    record[f"{q}_{metric}"]=item.get(metric)
            ordered=[by_date[d] for d in sorted(by_date)]
        else:
            ordered=raw
        for r in ordered:
            rows.append([s]+[r.get(h) for h in headers[1:]])
    fmts={h:FINANCE_FMT for h in headers if h not in {"scenario","date"}}
    write_table(ws,headers,rows,number_formats=fmts)
    style_sheet(ws,"B2")


REPRESENTATIVE_CASE_HEADERS = [
    "scenario","case","record_type","trial_id","date","final_close_date",
    "Cash","Net_Equity","Active_Bikes","Owned_Transferred_Bikes","Pending_Claims",
    "Partner1_Reinvestment_Balance","Partner2_Reinvestment_Balance",
    "final_net_project_equity","final_cash","cumulative_project_profit",
    "termination_count","secondary_cycle_count","owned_bikes","held_assets",
    "total_operating_revenue","bad_debt","guarantee_recovered","payload_source"
]


def start_representative_sheet(wb):
    ws = wb.create_sheet("REPRESENTATIVE_CASES")
    write_table(ws, REPRESENTATIVE_CASE_HEADERS, [])
    return ws


def _append_representative_row(ws, row):
    headers = REPRESENTATIVE_CASE_HEADERS
    if len(row) != len(headers):
        raise RuntimeError(f"REPRESENTATIVE_ROW_WIDTH_MISMATCH: {len(row)} != {len(headers)}")
    ws.append([py_value(value) for value in row])
    row_num = ws.max_row
    money_fields = {
        "Cash", "Net_Equity", "Partner1_Reinvestment_Balance",
        "Partner2_Reinvestment_Balance", "final_net_project_equity",
        "final_cash", "cumulative_project_profit", "total_operating_revenue",
        "bad_debt", "guarantee_recovered",
    }
    for cell in ws[row_num]:
        cell.alignment = RIGHT
        header = headers[cell.column - 1]
        if header in money_fields:
            cell.number_format = FINANCE_FMT
        if isinstance(cell.value, date):
            cell.number_format = DATE_FMT
    return row_num


def append_representative_scenario(
    ws,
    scenario,
    representative_payloads,
    assignments,
    role_to_payload_labels,
    trial_rows,
    *,
    on_p50=None,
):
    """Validate and write one scenario's payloads once; reuse P50 while building its detail sheets."""
    stage13 = _stage13_module()
    required_roles = stage13.REQUIRED_REPRESENTATIVE_ROLES
    summary_fields = (
        "final_net_project_equity", "final_cash", "cumulative_project_profit",
        "termination_count", "secondary_cycle_count", "owned_bikes", "held_assets",
        "total_operating_revenue", "bad_debt", "guarantee_recovered",
    )
    by_id = {}
    if trial_rows is not None:
        for trial in trial_rows:
            trial_id = trial.get("trial_id")
            if type(trial_id) is int:
                if trial_id in by_id:
                    raise RuntimeError(f"REPRESENTATIVE_TRIAL_ID_DUPLICATE: {scenario}/{trial_id}")
                by_id[trial_id] = trial

    headers = REPRESENTATIVE_CASE_HEADERS
    for role in required_roles:
        trial_id = assignments[role]
        if role == "MaxCase":
            row = [scenario, "MaxCase", "PENDING_DATA_SOURCE", trial_id]
            row.extend(["PENDING_DATA_SOURCE"] * (len(headers) - len(row)))
            row_num = _append_representative_row(ws, row)
            comment = (
                "MaxCase trial_id=60; daily payload missing from V3 — see R4-V2.1 §5.4"
                if scenario == "C050_G000"
                else f"MaxCase trial_id={trial_id}; daily payload missing from enrichment artifact — see R4-V2.1 §5.4"
            )
            ws.cell(row=row_num, column=1).comment = Comment(comment, "OpenAI")
            ws.cell(row=row_num, column=headers.index("Partner1_Reinvestment_Balance") + 1).comment = Comment(
                "PENDING_DATA_SOURCE: MaxCase partner 1 reinvestment balance absent from source", "OpenAI"
            )
            ws.cell(row=row_num, column=headers.index("Partner2_Reinvestment_Balance") + 1).comment = Comment(
                "PENDING_DATA_SOURCE: MaxCase partner 2 reinvestment balance absent from source", "OpenAI"
            )
            continue

        payload = representative_payloads[role]
        stage13.validate_representative_payload(scenario, role, trial_id, payload)
        summary = payload["summary"]
        if trial_rows is not None:
            source_row = by_id.get(trial_id)
            if source_row is None:
                raise RuntimeError(f"STAGE13_REPRESENTATIVE_TRIAL_BINDING_INVALID: {scenario}/{role}/trial_id={trial_id}")
            for field in summary_fields:
                if field in summary and field in source_row and summary[field] != source_row[field]:
                    raise RuntimeError(
                        f"STAGE13_REPRESENTATIVE_RESULT_MISMATCH: {scenario}/{role}/trial_id={trial_id}/{field}"
                    )
        source_ref = f"representatives/{role_to_payload_labels[role]}.json"
        _append_representative_row(ws, [
            scenario, role, "SUMMARY", trial_id, None, summary.get("final_close_date"),
            None, None, None, None, None, None, None,
            summary.get("final_net_project_equity"), summary.get("final_cash"),
            summary.get("cumulative_project_profit"), summary.get("termination_count"),
            summary.get("secondary_cycle_count"), summary.get("owned_bikes"), summary.get("held_assets"),
            summary.get("total_operating_revenue"), summary.get("bad_debt"), summary.get("guarantee_recovered"),
            source_ref
        ])
        for daily in payload.get("daily", []):
            _append_representative_row(ws, [
                scenario, role, "DAILY", trial_id, daily.get("date"),
                daily.get("Final_Close_Date"), daily.get("Cash"), daily.get("Net_Equity"),
                daily.get("Active_Bikes"), daily.get("Owned_Transferred_Bikes"),
                daily.get("Pending_Claims"), daily.get("Partner1_Reinvestment_Balance"),
                daily.get("Partner2_Reinvestment_Balance"), None, None, None, None, None,
                None, None, None, None, None, source_ref
            ])
        if role == "P50" and on_p50 is not None:
            on_p50(payload)
        del payload


def finish_representative_sheet(ws):
    style_sheet(ws, "E2")


def write_representatives(
    wb,
    manifest,
    representatives_by_scenario,
    assignments,
    role_to_payload_labels,
    trial_rows_by_scenario=None,
):
    """Non-incremental compatibility helper used by synthetic tests and legacy mode."""
    ws = start_representative_sheet(wb)
    for scenario in EXPECTED_SCENARIOS:
        rows = None if trial_rows_by_scenario is None else trial_rows_by_scenario[scenario]
        append_representative_scenario(
            ws, scenario, representatives_by_scenario[scenario], assignments[scenario],
            role_to_payload_labels, rows,
        )
    finish_representative_sheet(ws)




def write_stats_sheet(wb, scenario, rows, total_capital):
    ws=wb.create_sheet(f"{scenario}_STATS")
    headers=["metric","Mean","Median","Std","Min","Max","P5","P10","P25","P50","P75","P90","P95"]
    out=[]
    for label,field in numeric_metrics(rows):
        out.append([label,*metric_stats(rows,field)])
    profits=np.asarray([r["cumulative_project_profit"] for r in rows],dtype=float)
    out.append(["ROI",*(metric_stats([{"roi": p / total_capital} for p in profits],"roi"))])
    write_table(ws,headers,out,number_formats={h:FINANCE_FMT for h in headers[1:]})
    for row in ws.iter_rows(min_row=2):
        if row[0].value == "ROI":
            for cell in row[1:]:
                cell.number_format = PERCENT_FMT
    style_sheet(ws)


def write_daily_sheet(wb, scenario, payload):
    ws=wb.create_sheet(f"{scenario}_DAILY")
    rows=payload["daily"]
    headers=list(rows[0]) if rows else ["date"]
    # Materialize mappings as values in the exact worksheet header order.
    value_rows = [[row.get(header) for header in headers] for row in rows]
    write_table(
        ws,
        headers,
        value_rows,
        number_formats={
            h: FINANCE_FMT
            for h in headers
            if h not in {"date", "Final_Close_Date"}
        },
    )
    style_sheet(ws,"B2")


def write_assets_sheet(wb, scenario, payload):
    ws=wb.create_sheet(f"{scenario}_ASSETS")
    rows=payload.get("final_inventory",[])
    headers=["bike_id","source","purchase_date","delivery_date","gross_cost","accumulated_depreciation","net_book_value","termination_count","secondary_cycle_count","state"]
    write_table(ws,headers,[[r.get(h) for h in headers] for r in rows],number_formats={"gross_cost":FINANCE_FMT,"accumulated_depreciation":FINANCE_FMT,"net_book_value":FINANCE_FMT})
    style_sheet(ws)


def write_acctg_sheet(wb, scenario, payload):
    ws=wb.create_sheet(f"{scenario}_ACCTG")
    daily=payload["daily"]
    headers=["date","Operating_Revenue","Operating_Expenses","Operating_Net_Profit","Cumulative_Project_Profit","Cash","Accounts_Receivable","Guarantee_Claim_Receivable","Gross_Bike_Assets","Accumulated_Depreciation","Net_Equity"]
    rows=[[r.get(h) if h!="Net_Equity" else r.get("Net_Equity") for h in headers] for r in daily]
    write_table(ws,headers,rows,number_formats={h:FINANCE_FMT for h in headers if h!="date"})
    balance_checks=payload["accounting"].get("daily_balance_checks",[])
    start=ws.max_row+3
    if balance_checks:
        ws.cell(start,1,"Daily Balance Checks").font=Font(bold=True)
        bh=list(balance_checks[0])
        write_table(ws,bh,[[r.get(x) for x in bh] for r in balance_checks],start_row=start+1,number_formats={x:FINANCE_FMT for x in bh})
        start=ws.max_row+3
    for title,key in [("Cash Roll-forward","cash_rollforward"),("AR Roll-forward","ar_rollforward"),("Asset Roll-forward","asset_rollforward"),("Equity Roll-forward","equity_rollforward")]:
        ws.cell(start,1,title).font=Font(bold=True)
        data=payload["accounting"].get(key,[])
        if data:
            h=list(data[0])
            write_table(ws,h,[[r.get(x) for x in h] for r in data],start_row=start+1,number_formats={x:FINANCE_FMT for x in h})
            start=ws.max_row+3
    style_sheet(ws,"B2")


def write_partner_sheet(wb, scenario, payload):
    ws=wb.create_sheet(f"{scenario}_PARTNER")
    rows=payload.get("partner_memo",[])
    headers=["date","Partner1_Reinvestment_Balance","Partner2_Reinvestment_Balance"]
    write_table(ws,headers,[[r.get(h) for h in headers] for r in rows],number_formats={"Partner1_Reinvestment_Balance":FINANCE_FMT,"Partner2_Reinvestment_Balance":FINANCE_FMT})
    final=payload["final_reconciliation"]
    ws.cell(ws.max_row+3,1,"Partner1_Final_Entitlement").font=Font(bold=True)
    ws.cell(ws.max_row,2,payload["summary"]["partner1_final_entitlement"]).number_format=FINANCE_FMT
    ws.cell(ws.max_row+1,1,"Partner2_Final_Entitlement").font=Font(bold=True)
    ws.cell(ws.max_row,2,payload["summary"]["partner2_final_entitlement"]).number_format=FINANCE_FMT
    style_sheet(ws,"B2")


def write_generic(wb, title, sheet_name, items):
    ws=wb.create_sheet(sheet_name)
    if not items:
        write_table(ws,["message"],[["No records"]])
        style_sheet(ws)
        return
    if isinstance(items[0],dict):
        headers=list(items[0])
        write_table(ws,headers,[[x.get(h) for h in headers] for x in items],number_formats={h:FINANCE_FMT for h in headers})
    style_sheet(ws)


def write_scenario_detail(wb, scenario, payload):
    write_daily_sheet(wb,scenario,payload)
    write_assets_sheet(wb,scenario,payload)
    write_acctg_sheet(wb,scenario,payload)
    write_partner_sheet(wb,scenario,payload)
    write_generic(wb,"","".join([scenario,"_GUARANTEE"]),payload.get("guarantee_claims",[]))
    write_generic(wb,"","".join([scenario,"_CONTRACTS"]),payload.get("contracts",[]))


def collect_excel_errors(path: Path):
    wb=load_workbook(path,data_only=False)
    errors=[]
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.data_type=="e":
                    errors.append(f"{ws.title}!{cell.coordinate}")
    return errors


# The Stage 9C gate compares five numeric statistics using exact equality.
# openpyxl safe_string emits "%.16g", which can collapse a binary float such as
# 13185720.000000002 to 13185720. These helpers patch only affected gate cells.
STATS_PRECISION_METRICS = {
    "Final_Net_Project_Equity",
    "Partner1_Final_Entitlement",
    "Partner2_Final_Entitlement",
    "Final_Cash",
}
STATS_PRECISION_COLUMNS = (5, 6, 8, 10, 12)  # MIN, MAX, P10, P50, P90


def patch_xlsx_numeric_values(output_path: str | Path, targets: dict[str, dict[str, int | float]]) -> dict:
    """Patch selected numeric OOXML cells to exact round-trip decimal literals.

    targets maps worksheet name -> cell reference -> expected numeric value.
    Cells already equal as Python numbers are left unchanged. Every requested
    cell is verified to exist exactly once.
    """
    output_path = Path(output_path)
    namespaces = {
        "main": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
        "office_rel": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
        "package_rel": "http://schemas.openxmlformats.org/package/2006/relationships",
    }
    rel_id_attr = "{" + namespaces["office_rel"] + "}id"
    temp_fd, temp_name = tempfile.mkstemp(
        prefix=output_path.stem + ".precision-",
        suffix=output_path.suffix,
        dir=str(output_path.parent),
    )
    os.close(temp_fd)
    checked = 0
    patched: list[str] = []
    by_xml_path: dict[str, list[tuple[str, int | float]]] = {}
    try:
        with zipfile.ZipFile(output_path, "r") as source:
            wb_root = ET.fromstring(source.read("xl/workbook.xml"))
            rel_root = ET.fromstring(source.read("xl/_rels/workbook.xml.rels"))
            relationships = {
                rel.attrib["Id"]: rel.attrib["Target"]
                for rel in rel_root.findall("package_rel:Relationship", namespaces)
            }
            sheet_paths = {}
            for sheet in wb_root.findall("main:sheets/main:sheet", namespaces):
                sheet_name = sheet.attrib["name"]
                rel_id = sheet.attrib[rel_id_attr]
                target = relationships[rel_id].replace("\\", "/")
                if target.startswith("/"):
                    xml_path = target.lstrip("/")
                elif target.startswith("xl/"):
                    xml_path = target
                else:
                    xml_path = posixpath.normpath(posixpath.join("xl", target))
                sheet_paths[sheet_name] = xml_path

            for sheet_name, cells in targets.items():
                if sheet_name not in sheet_paths:
                    raise RuntimeError(f"PRECISION_SHEET_MISSING: {sheet_name}")
                xml_path = sheet_paths[sheet_name]
                for cell_ref, value in cells.items():
                    if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(float(value)):
                        raise RuntimeError(f"PRECISION_TARGET_NOT_FINITE_NUMERIC: {sheet_name}!{cell_ref}={value!r}")
                    by_xml_path.setdefault(xml_path, []).append((cell_ref, value))

            with zipfile.ZipFile(temp_name, "w") as dest:
                for item in source.infolist():
                    payload = source.read(item.filename)
                    for cell_ref, expected in by_xml_path.get(item.filename, []):
                        cell_pattern = re.compile(
                            rb'(<c\b(?=[^>]*\br="' + re.escape(cell_ref.encode("ascii")) +
                            rb'")[^>]*>.*?<v>)(.*?)(</v>.*?</c>)',
                            flags=re.DOTALL,
                        )
                        match = cell_pattern.search(payload)
                        checked += 1
                        if match is None:
                            raise RuntimeError(f"PRECISION_CELL_MISSING: {item.filename}/{cell_ref}")
                        lexical = match.group(2).decode("ascii")
                        try:
                            current = float(lexical)
                        except ValueError as exc:
                            raise RuntimeError(
                                f"PRECISION_CELL_NOT_NUMERIC: {item.filename}/{cell_ref}={lexical!r}"
                            ) from exc
                        if current != expected:
                            exact_lexical = repr(expected) if isinstance(expected, float) else str(expected)
                            payload, replacements = cell_pattern.subn(
                                lambda m: m.group(1) + exact_lexical.encode("ascii") + m.group(3),
                                payload,
                                count=1,
                            )
                            if replacements != 1:
                                raise RuntimeError(
                                    f"PRECISION_CELL_PATCH_COUNT: {item.filename}/{cell_ref}={replacements}"
                                )
                            patched.append(f"{item.filename}!{cell_ref}")
                    dest.writestr(item, payload)
        os.replace(temp_name, output_path)
    except Exception:
        try:
            os.unlink(temp_name)
        except FileNotFoundError:
            pass
        raise
    return {
        "target_cells_verified": checked,
        "cells_patched": len(patched),
        "patched_cells": patched,
        "numeric_cells_only": True,
        "uses_tolerance_or_rounding": False,
    }


def save_workbook_preserving_stats_precision(workbook: Workbook, output_path: str | Path) -> dict:
    """Save a workbook, restoring only gate cells changed by openpyxl serialization."""
    targets: dict[str, dict[str, int | float]] = {}
    for ws in workbook.worksheets:
        if not ws.title.endswith("_STATS"):
            continue
        for row_idx in range(2, ws.max_row + 1):
            metric = ws.cell(row_idx, 1).value
            if metric not in STATS_PRECISION_METRICS:
                continue
            for col_idx in STATS_PRECISION_COLUMNS:
                cell = ws.cell(row_idx, col_idx)
                value = cell.value
                if isinstance(value, float) and math.isfinite(value):
                    serialized = float("%.16g" % value)
                    if serialized != value:
                        targets.setdefault(ws.title, {})[cell.coordinate] = value
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return patch_xlsx_numeric_values(output_path, targets) if targets else {
        "target_cells_verified": 0,
        "cells_patched": 0,
        "patched_cells": [],
        "numeric_cells_only": True,
        "uses_tolerance_or_rounding": False,
    }


class _LazyScenarioRolePayloads:
    def __init__(self, root: Path, scenario: str, role_map: dict[str, str], assignments: dict[str, dict[str, int]]):
        self.root = root
        self.scenario = scenario
        self.role_map = role_map
        self.assignments = assignments

    def __getitem__(self, role: str):
        REQUIRED = _stage13_module().REQUIRED_REPRESENTATIVE_ROLES
        if role not in REQUIRED:
            raise KeyError(role)
        if role == "MaxCase":
            return {
                "_pending_data_source": True,
                "summary": {"scenario_id": self.scenario, "trial_id": self.assignments[self.scenario][role]},
                "daily": [],
            }
        return load_rep_payload(self.root, self.scenario, self.role_map[role])

    def __iter__(self):
        return iter(_stage13_module().REQUIRED_REPRESENTATIVE_ROLES)

    def __len__(self):
        return len(_stage13_module().REQUIRED_REPRESENTATIVE_ROLES)


class _LazyRepresentativesByScenario:
    def __init__(self, root: Path, role_map: dict[str, str], assignments: dict[str, dict[str, int]]):
        self.root = root
        self.role_map = role_map
        self.assignments = assignments

    def __getitem__(self, scenario: str):
        if scenario not in self.assignments:
            raise KeyError(scenario)
        return _LazyScenarioRolePayloads(self.root, scenario, self.role_map, self.assignments)

    def __iter__(self):
        return iter(self.assignments)

    def __len__(self):
        return len(self.assignments)


def build_workbook(
    dataset_root: str,
    output_path: str,
    manifest_path: str,
    representatives_path: str,
    *,
    role_to_payload_labels: dict[str, str] | None = None,
) -> dict:
    """Build a provisional workbook from pinned artifacts with single-pass representative reads."""
    stage13 = _stage13_module()
    validate_representative_assignments = stage13.validate_representative_assignments
    validate_role_payload_mapping = stage13.validate_role_payload_mapping

    root = Path(dataset_root)
    manifest = load_json(Path(manifest_path))
    manifest["_dataset_root"] = str(root)
    if manifest.get("source_sha") != EXPECTED_SOURCE_SHA:
        raise RuntimeError("SOURCE_SHA_MISMATCH")
    if manifest.get("scenarios") is None or list(sorted(manifest["scenarios"])) != list(sorted(EXPECTED_SCENARIOS)):
        raise RuntimeError("SCENARIO_IDENTITY_MISMATCH")

    assignment_object = load_json(Path(representatives_path))
    assignments = assignment_object.get("scenarios")
    if not isinstance(assignments, dict):
        raise ValueError("STAGE13_REPRESENTATIVE_ASSIGNMENTS_MISSING")
    validate_representative_assignments(assignments, expected_scenarios=EXPECTED_SCENARIOS)
    role_map = validate_role_payload_mapping(role_to_payload_labels)

    validation = (
        load_json(root / "validation_163.json")
        if (root / "validation_163.json").exists()
        else {"status": "not-provided", "tests": []}
    )
    errors = load_json(root / "error_log.json") if (root / "error_log.json").exists() else []
    all_rows = {scenario: load_scenario_trials(root, scenario) for scenario in EXPECTED_SCENARIOS}
    validate_representative_assignments(
        assignments,
        expected_scenarios=EXPECTED_SCENARIOS,
        trial_rows_by_scenario=all_rows,
    )

    data = {
        "manifest": manifest,
        "trials_by_scenario": all_rows,
        "daily_distributions_by_scenario": {},
        "representatives_by_scenario": _LazyRepresentativesByScenario(root, role_map, assignments),
        "representative_trial_ids": assignments,
        "validation": validation,
        "errors": errors,
        "dataset_root": str(root),
        "pending_roles": ["MaxCase daily payload missing across enrichment artifacts"],
    }
    return stage13.assemble_workbook(
        builder=__import__("sys").modules[__name__],
        data=data,
        output_path=output_path,
        role_to_payload_labels=role_map,
        mode="provisional-p16",
    )




def main():
    parser=argparse.ArgumentParser()
    parser.add_argument("--dataset-root",required=True)
    parser.add_argument("--output-path",required=True)
    parser.add_argument("--manifest",required=True)
    parser.add_argument("--representatives",required=True)
    parser.add_argument(
        "--role-payload-labels-json",
        help="Optional explicit JSON object mapping all five approved roles to payload filenames; omission fails closed.",
    )
    args=parser.parse_args()
    role_payload_labels = (
        load_json(Path(args.role_payload_labels_json))
        if args.role_payload_labels_json
        else None
    )
    result=build_workbook(
        args.dataset_root,args.output_path,args.manifest,args.representatives,
        role_to_payload_labels=role_payload_labels,
    )
    print(json.dumps(result,ensure_ascii=False,sort_keys=True))
    raise SystemExit(0 if result["success"] else 1)


if __name__=="__main__":
    main()
